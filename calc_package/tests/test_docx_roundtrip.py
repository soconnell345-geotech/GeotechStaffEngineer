"""Word -> Markdown -> Word: a memo survives being read and written again.

Live smoke wave 2a (B3, B4): a memo could not be read back faithfully (tables
came back one cell per line, list markers and line breaks were lost), so
every "edit the memo" rebuilt it from flat text and drifted; and write_docx
ran consecutive lines together, printed the title twice and signed the file
"python-docx". These tests read the documents back with python-docx and with
:func:`docx_to_markdown`, never by inspecting what was built.
"""

import os

import pytest

docx = pytest.importorskip("docx")
fitz = pytest.importorskip("fitz")
pytest.importorskip("markdown_it")

from calc_package.docx_renderer import (  # noqa: E402
    DEFAULT_AUTHOR, docx_to_markdown, markdown_to_docx,
)

MEMO = """\
**To:** Project file
**Re:** Summary of the submittal
**Status:** DRAFT

## Findings

1. Bearing elevation is unconfirmed
2. Logs end at 30 ft
   - nested bullet one
   - nested bullet two
3. Third item

Text with `inline code`, *italic* and **bold** words.

- bullet a
- bullet b

| Boring | Depth (ft) | Note |
| --- | --- | --- |
| B-1 | 30 | ok |
| B-2 | 25 | line one<br>line two |

*Table 1. Borings*

> A quoted remark

```
code line 1
code line 2
```

---

![Profile figure](profile.png)

1. a second list
2. numbered from one
"""


def _png(path):
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 16, 16), False)
    pix.set_rect(pix.irect, (200, 120, 40))
    path.write_bytes(pix.tobytes("png"))


@pytest.fixture
def memo(tmp_path):
    _png(tmp_path / "profile.png")
    path = markdown_to_docx(MEMO, str(tmp_path / "memo.docx"),
                            base_dir=str(tmp_path),
                            title="Review Memo - Project Alpha")
    return path, tmp_path


def test_a_memo_reads_back_as_the_markdown_it_was_written_from(memo):
    path, tmp = memo
    read = docx_to_markdown(path, image_dir=str(tmp / "img"), image_ref="img",
                            image_prefix="memo_")
    md = read["markdown"]
    assert read["title"] == "Review Memo - Project Alpha"
    assert md.startswith("# Review Memo - Project Alpha\n\n")
    # Line breaks kept, labels bold.
    assert "**To:** Project file\n**Re:** Summary of the submittal\n" \
           "**Status:** DRAFT" in md
    # Numbered list with a nested bulleted one, then the numbering resumes.
    assert ("1. Bearing elevation is unconfirmed\n2. Logs end at 30 ft\n"
            "   - nested bullet one\n   - nested bullet two\n3. Third item"
            ) in md
    assert "- bullet a\n- bullet b" in md
    assert "`inline code`" in md and "*italic*" in md and "**bold**" in md
    # A pipe table, its in-cell line break, and its caption.
    assert "| Boring | Depth (ft) | Note |\n| --- | --- | --- |" in md
    assert "| B-2 | 25 | line one<br>line two |" in md
    assert "*Table 1. Borings*" in md
    assert "> A quoted remark" in md
    assert "```\ncode line 1\ncode line 2\n```" in md
    assert "\n---\n" in md
    # The picture is saved beside and referenced, its caption as the alt.
    assert "![Profile figure](img/memo_fig1.png)" in md
    assert os.path.isfile(tmp / "img" / "memo_fig1.png")
    assert read["images"] == ["img/memo_fig1.png"]
    # A second numbered list starts at one again.
    assert "1. a second list\n2. numbered from one" in md


def test_read_edit_write_reaches_a_fixed_point(memo):
    """Writing the read-back (with its title) and reading again gives the
    same Markdown: a revision changes only what the editor changed."""
    path, tmp = memo
    first = docx_to_markdown(path, image_dir=str(tmp), image_ref=".",
                             image_prefix="rt_")
    again = markdown_to_docx(first["markdown"], str(tmp / "memo_v2.docx"),
                             base_dir=str(tmp), title=first["title"])
    second = docx_to_markdown(again, image_dir=str(tmp), image_ref=".",
                              image_prefix="rt_")
    assert second["markdown"] == first["markdown"]
    styles = [p.style.name for p in docx.Document(again).paragraphs]
    assert styles.count("Title") == 1
    assert "Heading 1" not in styles        # the title was not repeated


def test_consecutive_lines_stay_on_their_own_lines(tmp_path):
    path = markdown_to_docx("Project 1234, June 2026\nStatus: DRAFT\n",
                            str(tmp_path / "a.docx"))
    pars = [p for p in docx.Document(path).paragraphs if p.text.strip()]
    assert len(pars) == 1
    assert pars[0].text == "Project 1234, June 2026\nStatus: DRAFT"


@pytest.mark.parametrize("title, opening, kept", [
    # The title's own text, case and whitespace aside: written once (B4).
    ("Review Comments (DRAFT)", "#  review   COMMENTS (draft) ", False),
    ("Submittal Memo", "# **Submittal Memo**", False),
    # Anything more or less than the title is kept, so no word is lost
    # (live smoke 2c, E8: F44's longer heading was dropped with its figure
    # number, as were shorter ones that were the title's leading words).
    ("Review: Std. No. 21.01 Rev. 2, Bioretention Cross-Section",
     "# Review: Std. No. 21.01 Rev. 2, Bioretention Cross-Section "
     "(BMP Fig. 4.1.3)", True),
    ("Submittal Review Stamp - Project Alpha", "# Submittal Review Stamp",
     True),
    ("Review Comments (DRAFT)", "# Review comments - draft", True),
    # A one-word section is a section, not the title.
    ("Findings and Recommendations", "# Findings", True),
    # A different heading under the title stays.
    ("Foundation Review", "# Findings", True),
])
def test_the_title_is_written_once(tmp_path, title, opening, kept):
    path = markdown_to_docx(f"{opening}\n\nBody text.\n",
                            str(tmp_path / "t.docx"), title=title)
    pars = [(p.style.name, p.text) for p in docx.Document(path).paragraphs]
    assert pars[0] == ("Title", title)
    assert any(s == "Heading 1" for s, _t in pars) is kept
    if kept:            # with every one of its words
        heading = next(t for s, t in pars if s == "Heading 1")
        assert heading.split() == opening.lstrip("# ").split()


def test_file_info_names_the_author_and_the_title(tmp_path):
    plain = docx.Document(markdown_to_docx(
        "# Memo\n\nText.\n", str(tmp_path / "a.docx")))
    assert plain.core_properties.author == DEFAULT_AUTHOR \
        == "GeotechStaffEngineer"
    assert plain.core_properties.last_modified_by == DEFAULT_AUTHOR
    assert plain.core_properties.title == "Memo"     # the first heading
    signed = docx.Document(markdown_to_docx(
        "Text.\n", str(tmp_path / "b.docx"), title="Review",
        author="jdoe via GeotechStaffEngineer (AI draft)"))
    assert signed.core_properties.author == \
        "jdoe via GeotechStaffEngineer (AI draft)"
    assert signed.core_properties.title == "Review"
    assert "python-docx" not in (signed.core_properties.author,
                                 signed.core_properties.last_modified_by)


def test_each_numbered_list_restarts_at_its_own_first_number(tmp_path):
    path = markdown_to_docx("1. a\n2. b\n\nText.\n\n3. c\n4. d\n",
                            str(tmp_path / "n.docx"))
    doc = docx.Document(path)
    items = [p for p in doc.paragraphs if p.style.name == "List Number"]
    ids = [p._p.pPr.numPr.numId.val for p in items]
    assert ids[0] == ids[1] and ids[2] == ids[3] and ids[0] != ids[2]
    numbering = doc.part.numbering_part.element
    starts = {n.numId: [o.startOverride.val for o in n.lvlOverride_lst]
              for n in numbering.num_lst}
    assert starts[ids[0]] == [1] and starts[ids[2]] == [3]


def test_a_word_file_not_written_here_reads_too(tmp_path):
    """An uploaded memo: Word's own heading, list and table styles, a
    manual line break, a tracked insertion and deletion, a hyperlink-free
    body — and text that only LOOKS like Markdown stays text."""
    doc = docx.Document()
    doc.add_heading("Site Memo", 1)
    p = doc.add_paragraph()
    p.add_run("Prepared by: ").bold = True
    p.add_run("J. Smith")
    p.add_run().add_break()
    p.add_run("Date: 2026-10-09")
    doc.add_paragraph("- typed dash, not a list")
    doc.add_paragraph("first", style="List Bullet")
    doc.add_paragraph("second", style="List Number")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Value | unit"
    t.cell(1, 0).text, t.cell(1, 1).text = "Depth", "3.2 m"
    # A tracked change: "kept " inserted, "gone " deleted.
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    doc.element.body.insert(len(doc.element.body) - 1, parse_xml(
        f'<w:p {nsdecls("w")}><w:r><w:t xml:space="preserve">Text </w:t>'
        f'</w:r><w:ins w:id="1" w:author="x"><w:r><w:t xml:space="preserve">'
        f'kept </w:t></w:r></w:ins><w:del w:id="2" w:author="x"><w:r>'
        f'<w:delText>gone </w:delText></w:r></w:del><w:r><w:t>end</w:t>'
        f'</w:r></w:p>'))
    path = str(tmp_path / "upload.docx")
    doc.save(path)
    md = docx_to_markdown(path)["markdown"]
    assert md.startswith("# Site Memo\n")
    assert "**Prepared by:** J. Smith\nDate: 2026-10-09" in md
    assert "\\- typed dash, not a list" in md
    assert "- first" in md and "1. second" in md
    assert "| Item | Value \\| unit |" in md and "| Depth | 3.2 m |" in md
    assert "Text kept end" in md and "gone" not in md
