"""The hand-labelled page types, mapped onto the ingest's label vocabulary.

The ground truth is a spreadsheet in the gitignored corpus folder: 4,300 pages
of 15 reports, labelled by hand. Its first sheet is a summary that names each
label sheet's report BY FILE NAME, and every other sheet has one row per page
with the page's text and its ``page_type``.

PRIVACY. Sheet names, report names and engineer names in that workbook are
private and must never reach a tracked file -- including the sheet names, which
carry a firm and a city. The sheet-to-ID matching is therefore derived at
RUNTIME from the manifest's private stems and cached into
``raw/labels_map.json``; everything this module returns or prints is an ID, a
sheet index, a page count or a label.

:data:`RAW_TO_LABEL` is where the vocabularies meet. It is deliberately total
over the 23 raw strings the spreadsheet actually uses: an unknown string raises
rather than being dropped, because a silently ignored label is a silently wrong
score.
"""

from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple

from module_work.report_ingest_harness.corpus import (
    LABELS_MAP_JSON, LABELS_XLSX, list_reports, raw_available, report,
)

#: The page-label vocabulary (plan section 4, plus ``field_test`` and
#: ``letter``, which the hand labels distinguish and the plan's list did not).
LABELS: Tuple[str, ...] = (
    "narrative",
    "figure",
    "plan",
    "profile",
    "boring_log",
    "test_pit_log",
    "cpt_log",
    "dcp_log",
    "lab_test",
    "field_test",
    "calculation",
    "appended_report",
    "photos",
    "divider",
    "cover",
    "letter",
    "toc",
    "other",
)

#: Every ``page_type`` string the spreadsheet uses (lower-cased, stripped),
#: mapped into :data:`LABELS`. Change a row here, not in the spreadsheet.
RAW_TO_LABEL: Dict[str, str] = {
    "main report narrative": "narrative",
    "boring log": "boring_log",
    "test pit log": "test_pit_log",
    "cpt log": "cpt_log",
    "dcp log": "dcp_log",
    "lab testing": "lab_test",
    "subsurface drainage testing": "field_test",
    "calculation": "calculation",
    "appended report": "appended_report",
    "field photos": "photos",
    "appendix cover page": "divider",
    "figures or tables cover page": "divider",
    "figures or tables cover sheet": "divider",
    "figures or reports cover page": "divider",
    "cover page": "cover",
    "cover letter": "letter",
    "table of contents": "toc",
    "figure": "figure",
    "test location plan": "plan",
    "subsurface profile": "profile",
    "informational appendix content": "other",
    "other or unknown": "other",
    "other appendix table": "other",
}


class UnknownPageType(ValueError):
    """A ``page_type`` the mapping does not cover."""


def to_label(raw: str) -> str:
    """Map one raw ``page_type`` string into :data:`LABELS`."""
    key = " ".join(str(raw).strip().lower().split())
    try:
        return RAW_TO_LABEL[key]
    except KeyError:
        raise UnknownPageType(
            f"page_type {key!r} is not in RAW_TO_LABEL; add a row rather than "
            f"letting the page fall out of the score") from None


@dataclass(frozen=True)
class PageLabel:
    """One hand-labelled page. ``page0`` is 0-based; the sheet is 1-based."""

    page0: int
    raw: str
    label: str


@dataclass(frozen=True)
class SheetMatch:
    """How one label sheet was matched to a corpus ID.

    Safe to print in full: an index, a report ID and two page counts. The
    sheet's own name is deliberately not a field.
    """

    index: int
    rows: int
    rid: Optional[str]
    method: str                      # "name", "pages" or "unmatched"
    manifest_pages: Optional[int] = None

    @property
    def rows_match(self) -> Optional[bool]:
        if self.manifest_pages is None:
            return None
        return self.rows == self.manifest_pages

    def describe(self) -> str:
        if self.rid is None:
            return (f"sheet {self.index}: NO MATCH ({self.rows} label rows) "
                    f"-- no corpus PDF matches its report name, and nothing "
                    f"with that page count shares a word with it")
        verdict = "rows == pages" if self.rows_match else (
            f"MISMATCH: {self.rows} label rows vs {self.manifest_pages} pages")
        return (f"sheet {self.index}: {self.rid} "
                f"({self.rows} rows, matched by {self.method}) -- {verdict}")


# -- name normalisation -----------------------------------------------------

def _norm(name: str) -> str:
    """Case, punctuation and accent-insensitive form of a file stem."""
    text = unicodedata.normalize("NFKD", str(name))
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = text.lower()
    if text.endswith(".pdf"):
        text = text[:-4]
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def _tokens(name: str) -> List[str]:
    """Distinctive words of a file stem: three characters or more."""
    return [t for t in _norm(name).split() if len(t) >= 3]


def _workbook():
    try:
        import openpyxl
    except ImportError as exc:                       # pragma: no cover
        raise ImportError(
            "reading the hand labels needs openpyxl (it is in the app's "
            "dev venv)") from exc
    if not LABELS_XLSX.is_file():
        raise FileNotFoundError(
            f"the hand-label spreadsheet is not on this machine "
            f"({LABELS_XLSX})")
    return openpyxl.load_workbook(LABELS_XLSX, read_only=True, data_only=True)


def _sheet_rows(ws) -> List[dict]:
    """Rows of a label sheet as dicts keyed by the header, minus the header."""
    out: List[dict] = []
    header: Optional[List[str]] = None
    for row in ws.iter_rows(values_only=True):
        if header is None:
            header = [(" ".join(str(c).strip().lower().split()) if c else "")
                      for c in row]
            continue
        if all(c is None for c in row):
            continue
        out.append(dict(zip(header, row)))
    return out


# -- sheet -> ID matching ---------------------------------------------------

_CACHE: Optional[Tuple[List[SheetMatch], Dict[int, str]]] = None


def _build_map() -> Tuple[List[SheetMatch], Dict[int, str]]:
    wb = _workbook()
    names = list(wb.sheetnames)
    # The summary sheet: sheet name -> report file name.
    summary = {}
    for row in _sheet_rows(wb[names[0]]):
        sheet = row.get("sheet")
        rname = row.get("report name")
        if sheet and rname:
            summary[str(sheet)] = str(rname)

    by_norm: Dict[str, str] = {}
    for info in list_reports():
        by_norm.setdefault(_norm(info.private_name), info.id)

    matches: List[SheetMatch] = []
    sheet_of: Dict[int, str] = {}
    taken: set = set()
    pending: List[Tuple[int, str, int]] = []         # index, sheet name, rows

    for index, sheet_name in enumerate(names):
        if index == 0:
            continue
        rows = len(_sheet_rows(wb[sheet_name]))
        rid = by_norm.get(_norm(summary.get(sheet_name, "")))
        if rid is not None and rid not in taken:
            taken.add(rid)
            sheet_of[index] = sheet_name
            matches.append(SheetMatch(index, rows, rid, "name",
                                      report(rid).pages))
        else:
            pending.append((index, sheet_name, rows))

    # Fallback: a unique corpus report with exactly that many pages AND at
    # least one distinctive token in common with the sheet's report name.
    #
    # The token guard is not fussiness. Without it the one sheet whose report
    # was never copied into the corpus (153 rows) matched an unrelated
    # 153-page report on the page count alone, which would have scored 153
    # hand labels against the wrong document. A page count is a weak signal;
    # two weak signals that agree are worth acting on, one is not.
    for index, sheet_name, rows in pending:
        wanted = set(_tokens(summary.get(sheet_name, "")))
        candidates = [r for r in list_reports()
                      if r.pages == rows and r.id not in taken]
        shared = [r for r in candidates
                  if wanted & set(_tokens(r.private_name))]
        if len(candidates) == 1 and len(shared) == 1:
            rid = shared[0].id
            taken.add(rid)
            sheet_of[index] = sheet_name
            matches.append(SheetMatch(index, rows, rid, "pages",
                                      report(rid).pages))
        else:
            matches.append(SheetMatch(index, rows, None, "unmatched"))

    matches.sort(key=lambda m: m.index)
    wb.close()
    return matches, sheet_of


def sheet_map(refresh: bool = False) -> List[SheetMatch]:
    """Match every label sheet to a corpus ID, caching into ``raw/``.

    The cache holds the sheet NAMES (it lives in the gitignored corpus folder,
    which is where private strings belong); what this function returns holds
    none.
    """
    global _CACHE
    if _CACHE is not None and not refresh:
        return list(_CACHE[0])
    if not refresh and LABELS_MAP_JSON.is_file():
        try:
            blob = json.loads(LABELS_MAP_JSON.read_text(encoding="utf-8"))
            matches = [SheetMatch(**{k: v for k, v in row.items()
                                     if k != "sheet_name"})
                       for row in blob["sheets"]]
            sheet_of = {int(row["index"]): row["sheet_name"]
                        for row in blob["sheets"] if row.get("sheet_name")}
            _CACHE = (matches, sheet_of)
            return list(matches)
        except (OSError, ValueError, KeyError, TypeError):
            pass                                     # rebuild below
    matches, sheet_of = _build_map()
    _CACHE = (matches, sheet_of)
    try:
        LABELS_MAP_JSON.write_text(json.dumps({
            "note": "derived at runtime; PRIVATE (sheet names) -- "
                    "raw/ is gitignored",
            "sheets": [dict(index=m.index, rows=m.rows, rid=m.rid,
                            method=m.method,
                            manifest_pages=m.manifest_pages,
                            sheet_name=sheet_of.get(m.index))
                       for m in matches],
        }, indent=2), encoding="utf-8")
    except OSError:                                  # pragma: no cover
        pass
    return list(matches)


def mapped_ids() -> List[str]:
    """Corpus IDs that have hand labels, in ID order."""
    return sorted(m.rid for m in sheet_map() if m.rid)


def _sheet_name_for(rid: str) -> str:
    sheet_map()
    assert _CACHE is not None
    for index, name in _CACHE[1].items():
        for m in _CACHE[0]:
            if m.index == index and m.rid == rid:
                return name
    raise KeyError(f"{rid} has no hand-label sheet")


def labels_for(rid: str) -> List[PageLabel]:
    """The hand labels for one report, 0-based, in page order."""
    name = _sheet_name_for(rid)
    wb = _workbook()
    try:
        rows = _sheet_rows(wb[name])
    finally:
        wb.close()
    out: List[PageLabel] = []
    for row in rows:
        page = row.get("page")
        raw = row.get("page_type")
        if page is None or raw is None:
            continue
        out.append(PageLabel(page0=int(page) - 1,
                             raw=" ".join(str(raw).strip().lower().split()),
                             label=to_label(raw)))
    out.sort(key=lambda p: p.page0)
    return out


def label_counts() -> Dict[str, int]:
    """Hand-label counts per mapped label, over every matched report."""
    counts = {label: 0 for label in LABELS}
    for rid in mapped_ids():
        for pl in labels_for(rid):
            counts[pl.label] += 1
    return counts


def labels_available() -> bool:
    return raw_available() and LABELS_XLSX.is_file()
