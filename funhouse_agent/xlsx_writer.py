"""Excel workbooks from tables -- the ``write_xlsx`` tool's engine.

Live smoke wave 1 (A12): asked for a spreadsheet, the agent said "I could not
write a native .xlsx file with the tools I have" and saved SpreadsheetML as
``.xls``, which Excel opens only after a warning. Neither page could write
Excel. This module turns tables the agent already has -- rows, or the pipe
tables of a Markdown answer -- into a real ``.xlsx``, one sheet per table.

Pure function of its input (bytes out, no file I/O): the tool layer
(:mod:`funhouse_agent.deep.tools`) decides where the file goes, through the
host's save function, so it lands in the conversation's working folder and
gets a download card like every other saved file. openpyxl arrives with the
app (``python-ags4`` requires it); :func:`available` says whether it is here.
"""

from __future__ import annotations

import io
import re
from typing import Any, Dict, Iterable, List, Optional, Tuple

#: Characters Excel refuses in a sheet name, and its length limit.
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_SHEET_NAME = 31
#: Widest a column is sized to (characters). A longer cell WRAPS inside it
#: (live smoke 2c, E7: F32's issues ran to 272 characters in a 60-wide column
#: with no wrap, and showed as one clipped line).
MAX_COLUMN_WIDTH = 60
#: Narrowest a column is sized to (characters).
MIN_COLUMN_WIDTH = 8
#: The height of one line of Excel's default font (Calibri 11), and Excel's
#: tallest row, in points: a row holding wrapped cells is set tall enough
#: for its longest one, so every viewer shows the whole text (Excel alone
#: re-fits a row on opening; previews and other readers do not).
LINE_HEIGHT_PT = 15.0
MAX_ROW_HEIGHT_PT = 409.0
#: Most rows written to one sheet (Excel's own limit is 1,048,576).
MAX_ROWS = 100_000

_NUMBER = re.compile(r"^[+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?$")
_THOUSANDS = re.compile(r"^[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?$")
_EMPHASIS = re.compile(r"^(\*\*|__|\*|_|`)(.+)\1$")
_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.*?)\s*#*\s*$")
_SEPARATOR_CELL = re.compile(r"^:?-{2,}:?$")


def available() -> bool:
    """Whether openpyxl can be imported here."""
    try:
        import openpyxl  # noqa: F401
    except Exception:                                  # noqa: BLE001
        return False
    return True


# ---------------------------------------------------------------------------
# Markdown pipe tables
# ---------------------------------------------------------------------------

def _split_row(line: str) -> List[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|") and not text.endswith("\\|"):
        text = text[:-1]
    cells, cur, i = [], [], 0
    while i < len(text):
        ch = text[i]
        if ch == "\\" and i + 1 < len(text) and text[i + 1] == "|":
            cur.append("|")
            i += 2
            continue
        if ch == "|":
            cells.append("".join(cur).strip())
            cur = []
        else:
            cur.append(ch)
        i += 1
    cells.append("".join(cur).strip())
    return cells


def _is_separator(line: str) -> bool:
    cells = _split_row(line)
    return bool(cells) and all(_SEPARATOR_CELL.match(c.replace(" ", ""))
                               for c in cells if c != "") and \
        any(c for c in cells)


def tables_from_markdown(markdown: str) -> List[Dict[str, Any]]:
    """Every pipe table in ``markdown`` as ``{"name", "rows"}`` -- the header
    row first. A table is named by the nearest heading above it (``Table N``
    when there is none)."""
    lines = str(markdown or "").splitlines()
    tables: List[Dict[str, Any]] = []
    heading: Optional[str] = None
    i = 0
    while i < len(lines):
        line = lines[i]
        m = _HEADING.match(line)
        if m:
            heading = m.group(1).strip() or heading
            i += 1
            continue
        if line.strip().startswith("|") and i + 1 < len(lines) \
                and _is_separator(lines[i + 1]):
            rows = [_split_row(line)]
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(_split_row(lines[j]))
                j += 1
            tables.append({"name": heading or f"Table {len(tables) + 1}",
                           "rows": rows})
            heading = None
            i = j
            continue
        i += 1
    return tables


# ---------------------------------------------------------------------------
# Cells and sheets
# ---------------------------------------------------------------------------

def cell_value(value: Any) -> Any:
    """What goes in the cell: numbers as numbers (``12.5``, ``-3``,
    ``1,250``, ``4.2e-3``), Markdown emphasis taken off, everything else as
    text. A number written with a leading zero (an id such as ``007``) stays
    text; so does a percentage, a range or a value with a unit."""
    if value is None:
        return None
    if isinstance(value, bool) or isinstance(value, (int, float)):
        return value
    text = str(value).strip()
    m = _EMPHASIS.match(text)
    while m:
        text = m.group(2).strip()
        m = _EMPHASIS.match(text)
    if text == "":
        return None
    plain = text.replace(",", "") if _THOUSANDS.match(text) else text
    if _NUMBER.match(plain):
        digits = plain.lstrip("+-")
        if len(digits) > 1 and digits[0] == "0" and digits[1].isdigit():
            return text                    # an id, not a number
        try:
            number = float(plain)
        except ValueError:
            return text
        if re.match(r"^[+-]?\d+$", plain) and abs(number) < 2 ** 53:
            return int(plain)
        return number
    return text


def _rows_of(raw: Any) -> List[List[Any]]:
    """Rows as lists: lists/tuples as given; dicts become a header row of
    their keys (first-seen order) and a row of values each."""
    rows = list(raw or [])
    if rows and all(isinstance(r, dict) for r in rows):
        keys: List[str] = []
        for r in rows:
            for k in r:
                if k not in keys:
                    keys.append(k)
        return [list(keys)] + [[r.get(k) for k in keys] for r in rows]
    out = []
    for r in rows:
        if isinstance(r, (list, tuple)):
            out.append(list(r))
        elif isinstance(r, dict):
            out.append(list(r.values()))
        else:
            out.append([r])
    return out


def sheet_name(name: Any, taken: Iterable[str]) -> str:
    """An Excel-safe sheet name, unique (case-insensitive) among ``taken``."""
    base = _BAD_SHEET_CHARS.sub("-", str(name or "").strip()).strip("'")
    base = (base or "Sheet")[:MAX_SHEET_NAME]
    used = {t.lower() for t in taken}
    if base.lower() not in used:
        return base
    n = 2
    while True:
        tail = f" ({n})"
        cand = base[:MAX_SHEET_NAME - len(tail)] + tail
        if cand.lower() not in used:
            return cand
        n += 1


def _lines(value: Any) -> List[str]:
    return str(value).split("\n") if value is not None else []


def _longest_line(value: Any) -> int:
    """The longest line of a cell's text, in characters (0 when empty)."""
    return max((len(line) for line in _lines(value)), default=0)


def column_width(longest: int) -> float:
    """A column's width (Excel characters) for its longest line of text:
    that plus a margin, between :data:`MIN_COLUMN_WIDTH` and
    :data:`MAX_COLUMN_WIDTH` (a longer cell wraps inside it)."""
    return float(min(MAX_COLUMN_WIDTH, max(MIN_COLUMN_WIDTH, longest + 2)))


def _wrapped_lines(value: Any, width: float) -> int:
    """How many lines a cell's text takes when wrapped in a column
    ``width`` characters wide (generous: proportional letters are mostly
    narrower than the '0' Excel's width counts in)."""
    per_line = max(1, int(width) - 1)
    return sum(max(1, -(-len(line) // per_line)) for line in _lines(value))


def _wrap_and_fit(ws, width_of: Dict[int, float], Alignment) -> int:
    """Top-align every cell; wrap each cell longer than its column or
    holding a line break, and make its row tall enough to show it. Returns
    how many cells wrap."""
    n_wrapped = 0
    for row in ws.iter_rows():
        lines_needed = 1
        for cell in row:
            v = cell.value
            width = width_of.get(cell.column, MIN_COLUMN_WIDTH)
            wrap = isinstance(v, str) and (
                "\n" in v or _longest_line(v) > width - 2)
            cell.alignment = Alignment(vertical="top", wrap_text=wrap)
            if wrap:
                n_wrapped += 1
                lines_needed = max(lines_needed, _wrapped_lines(v, width))
        if lines_needed > 1:
            ws.row_dimensions[row[0].row].height = min(
                MAX_ROW_HEIGHT_PT, LINE_HEIGHT_PT * lines_needed)
    return n_wrapped


def build_workbook(sheets: Optional[List[Dict[str, Any]]] = None,
                   markdown: str = "",
                   header: bool = True) -> Tuple[bytes, List[Dict[str, Any]],
                                                 List[str]]:
    """``(xlsx_bytes, summary, warnings)``.

    ``sheets``: ``[{"name": str, "rows": [[...], ...] or [{...}, ...],
    "header": bool}]``; ``markdown``: text whose pipe tables become sheets
    after those. The first row of each sheet is its header (bold, frozen)
    unless ``header`` is false for it. Each column is sized to its longest
    line of text (:data:`MIN_COLUMN_WIDTH` to :data:`MAX_COLUMN_WIDTH`), a
    cell longer than its column (or holding line breaks) wraps, every cell
    is top-aligned, and a row with a wrapped cell is made tall enough to
    show it. Raises ``ValueError`` when there is no table at all."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    specs: List[Dict[str, Any]] = []
    for i, spec in enumerate(sheets or []):
        if isinstance(spec, dict):
            specs.append({"name": spec.get("name") or f"Sheet{i + 1}",
                          "rows": _rows_of(spec.get("rows")),
                          "header": bool(spec.get("header", header))})
        else:                                  # a bare list of rows
            specs.append({"name": f"Sheet{i + 1}", "rows": _rows_of(spec),
                          "header": header})
    for table in tables_from_markdown(markdown):
        specs.append({"name": table["name"], "rows": table["rows"],
                      "header": header})
    specs = [s for s in specs if s["rows"]]
    if not specs:
        raise ValueError(
            "no table to write: give 'sheets' as [{name, rows: [[header...], "
            "[row...], ...]}] or 'markdown' containing pipe tables "
            "(| a | b | with a |---|---| line under the header)")

    wb = Workbook()
    wb.remove(wb.active)
    summary: List[Dict[str, Any]] = []
    warnings: List[str] = []
    names: List[str] = []
    for spec in specs:
        name = sheet_name(spec["name"], names)
        names.append(name)
        ws = wb.create_sheet(title=name)
        rows = spec["rows"]
        if len(rows) > MAX_ROWS:
            warnings.append(f"sheet '{name}': {len(rows):,} rows given, the "
                            f"first {MAX_ROWS:,} written")
            rows = rows[:MAX_ROWS]
        widths: Dict[int, int] = {}
        n_cols = 0
        for r, row in enumerate(rows, start=1):
            n_cols = max(n_cols, len(row))
            for c, value in enumerate(row, start=1):
                v = cell_value(value)
                if v is not None and not isinstance(v, (int, float, str)):
                    v = str(v)
                cell = ws.cell(row=r, column=c, value=v)
                if spec["header"] and r == 1:
                    cell.font = Font(bold=True)
                widths[c] = max(widths.get(c, 0), _longest_line(v))
        width_of = {c: column_width(w) for c, w in widths.items()}
        for c, w in width_of.items():
            ws.column_dimensions[get_column_letter(c)].width = w
        n_wrapped = _wrap_and_fit(ws, width_of, Alignment)
        if spec["header"] and len(rows) > 1:
            ws.freeze_panes = "A2"
        summary.append({"sheet": name, "rows": len(rows), "columns": n_cols,
                        "header": spec["header"],
                        **({"wrapped_cells": n_wrapped} if n_wrapped
                           else {})})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), summary, warnings


__all__ = ["available", "build_workbook", "tables_from_markdown",
           "cell_value", "column_width", "sheet_name", "MAX_SHEET_NAME",
           "MAX_ROWS", "MAX_COLUMN_WIDTH"]
