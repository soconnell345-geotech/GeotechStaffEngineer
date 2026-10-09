"""Live smoke wave 2b: the document tools (module_work/live_smoke/runs/
w2b-review-new-sonnet/REVIEW.md).

* C2  a Word or Excel file goes to the Markdown reader, never to MuPDF
      (F41: a valid workbook read as "0 0 0 0 0 1 0 1 0").
* C8  a DXF is read as CAD data -- its text, leaders, dimensions, layers --
      and the page tools say what is possible instead of failing.
* C3  a document tool that raises returns a JSON error with no server path.
* C9  every page RANGE in a result carries the viewer's numbers too.
* C4  a marked copy is written to a temporary file and swapped in, even
      while this conversation holds it open; one mark is removed or replaced
      by its id or its words.

Offline: planlens' synthetic review document, files made here, no engine.
"""

from __future__ import annotations

import json
import os

import pytest

pytest.importorskip("planlens.tools")
fitz = pytest.importorskip("fitz")

from funhouse_agent import document_tools  # noqa: E402
from funhouse_agent.deep.tools import make_vision_tools  # noqa: E402
from planlens.testing import build_synthetic_review_document  # noqa: E402


@pytest.fixture(scope="module")
def gt():
    return build_synthetic_review_document()


@pytest.fixture
def folder(tmp_path, monkeypatch):
    """A conversation's working folder, bound as the host binds it."""
    work = tmp_path / "conv" / "files"
    work.mkdir(parents=True)
    monkeypatch.setenv("GEOTECH_DEFAULT_OUTPUT_DIR", str(work))
    monkeypatch.delenv("GEOTECH_MARKUP_AUTHOR", raising=False)
    return work


def _tool(tools, name):
    return next(t for t in tools if t.name == name)


def _invoke(tool, **kwargs):
    return json.loads(tool.invoke(kwargs))


# ---------------------------------------------------------------------------
# C2: Word and Excel
# ---------------------------------------------------------------------------

def _submittal_log(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Log"
    ws.append(["No.", "Item", "Spec section", "Status", "Rev"])
    rows = [("S-01", "Concrete mix 3000 psi", "32 16 00", "Submitted", 0),
            ("S-02", "Curb ramp detail", "32 16 00", "Resubmit", 1),
            ("S-03", "Underdrain pipe", "33 46 00", "Approved", 0)]
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def _spec_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("SECTION 32 16 00 - SITE CONCRETE", level=1)
    d.add_paragraph("A. Concrete: 3600 psi at 28 days.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Limit"
    t.cell(1, 0).text, t.cell(1, 1).text = "Cross slope", "2.0 %"
    d.save(path)


def test_an_excel_log_reads_as_its_table_not_nine_digits(folder):
    """F41: open_document + read_document on the log gave only the numeric
    Rev column; every text cell was missing."""
    _submittal_log(folder / "Submittal Log.xlsx")
    tools = make_vision_tools(engine=None)
    opened = _invoke(_tool(tools, "open_document"),
                     source="Submittal Log.xlsx")
    assert "error" not in opened, opened
    assert opened["kind"] == "excel workbook"
    assert opened["handle"] == "Submittal Log.xlsx"
    assert "Concrete mix 3000 psi" in opened["text"]
    assert "Resubmit" in opened["text"]
    assert "no pages" in opened["note"]
    read = _invoke(_tool(tools, "read_document"),
                   handle=opened["handle"], start_line=0)
    assert "Underdrain pipe" in read["text"]
    hits = _invoke(_tool(tools, "search_document"),
                   handle="Submittal Log.xlsx", pattern="resubmit")
    assert hits["n_hits"] == 1 and "S-02" in hits["hits"][0]["text"]


def test_a_word_spec_reads_as_markdown_and_pages_by_line(folder):
    _spec_docx(folder / "Section 32 16 00.docx")
    tools = make_vision_tools(engine=None)
    opened = _invoke(_tool(tools, "open_document"),
                     source="Section 32 16 00.docx")
    assert opened["kind"] == "word document"
    assert "# SECTION 32 16 00" in opened["text"]
    assert "| Cross slope | 2.0 % |" in opened["text"]
    # a page tool says what it cannot do, and what to use
    out = _invoke(_tool(tools, "document_page_map"),
                  handle="Section 32 16 00.docx")
    assert "no pages" in out["error"]
    assert "read_document" in out["hint"]


def test_a_long_office_file_pages_through_start_line(folder):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["#", "text"])
    for i in range(1500):
        ws.append([i, f"row {i} " + "lorem ipsum " * 6])
    wb.save(folder / "big.xlsx")
    tools = make_vision_tools(engine=None)
    first = _invoke(_tool(tools, "read_document"), handle="big.xlsx")
    assert "next" in first and first["next"]["start_line"] > 0
    second = _invoke(_tool(tools, "read_document"), handle="big.xlsx",
                     start_line=first["next"]["start_line"])
    assert second["text"] and second["text"] != first["text"]


def test_a_damaged_workbook_is_said_plainly(folder):
    (folder / "broken.xlsx").write_bytes(b"PK\x03\x04 not really a zip")
    tools = make_vision_tools(engine=None)
    out = _invoke(_tool(tools, "open_document"), source="broken.xlsx")
    assert "could not be read" in out["error"]
    assert str(folder) not in json.dumps(out)


def test_an_uploaded_workbook_held_only_as_bytes_is_read(tmp_path):
    path = tmp_path / "log.xlsx"
    _submittal_log(path)
    tools = make_vision_tools(engine=None,
                              attachments={"log.xlsx": path.read_bytes()})
    out = _invoke(_tool(tools, "open_document"), source="log.xlsx")
    assert "Curb ramp detail" in out["text"]


# ---------------------------------------------------------------------------
# C8: DXF
# ---------------------------------------------------------------------------

def _boring_plan(path):
    import ezdxf
    doc = ezdxf.new()
    doc.header["$INSUNITS"] = 2                         # feet
    msp = doc.modelspace()
    doc.layers.add("BORINGS")
    msp.add_line((0, 0), (400, 0), dxfattribs={"layer": "SITE"})
    msp.add_text("B-2", dxfattribs={"insert": (100, 200), "height": 5,
                                    "layer": "BORINGS"})
    msp.add_text("C-2 BORING PLAN", dxfattribs={"insert": (10, 380),
                                               "height": 8, "layer": "TITLE"})
    doc.saveas(path)


def test_a_dxf_is_read_as_cad_data_in_its_own_units(folder):
    """F40: open_document said "could not open ... as a PDF or image:
    FileDataError" and the model read raw DXF 6,000 characters a call."""
    _boring_plan(folder / "C-2_Boring_Plan.dxf")
    tools = make_vision_tools(engine=None)
    out = _invoke(_tool(tools, "open_document"),
                  source="C-2_Boring_Plan.dxf")
    assert "error" not in out, out
    assert out["kind"] == "dxf drawing"
    assert out["drawing_units"] == "ft"
    assert out["layers"]["BORINGS"] == 1
    assert 'TEXT "B-2" at (100, 200) [layer BORINGS]' in out["text"]
    # reading order: the title (y 380) before the boring (y 200)
    assert out["text"].index("BORING PLAN") < out["text"].index('"B-2"')
    assert "PDF plot" in out["note"]
    hits = _invoke(_tool(tools, "search_document"),
                   handle="C-2_Boring_Plan.dxf", pattern="B-2")
    assert hits["n_hits"] == 1
    refused = _invoke(_tool(tools, "annotate_document"),
                      handle="C-2_Boring_Plan.dxf",
                      markups=[{"kind": "note", "page": 0, "comment": "x",
                                "point": [1, 1]}])
    assert "error" in refused


# ---------------------------------------------------------------------------
# C3: a document tool that raises
# ---------------------------------------------------------------------------

def test_a_raising_tool_is_a_json_error_without_a_server_path(
        gt, folder, monkeypatch):
    from planlens.tools import ReviewToolkit

    def explode(self, name, arguments, max_chars=None):
        raise RuntimeError(f"code=2: cannot remove file '{folder}"
                           f"{os.sep}review_set_marked.pdf': Permission "
                           "denied")

    monkeypatch.setattr(ReviewToolkit, "call_json", explode)
    out = json.loads(document_tools.dispatch_document_tool(
        "open_document", {"source": "s.pdf"}, attachments={"s.pdf": gt.pdf}))
    assert out["error"].startswith("open_document failed: RuntimeError")
    assert "review_set_marked.pdf" in out["error"]
    assert str(folder) not in json.dumps(out)
    assert "hint" in out


# ---------------------------------------------------------------------------
# C9: page ranges in the viewer's numbers
# ---------------------------------------------------------------------------

def test_every_page_range_carries_the_viewers_numbers(gt, folder):
    tools = make_vision_tools(engine=None, attachments={"s.pdf": gt.pdf})
    out = _invoke(_tool(tools, "open_document"), source="s.pdf")
    by_kind = out["pages_by_kind"]
    shown = out["pdf_pages_by_kind"]
    for kind, ranges in by_kind.items():
        assert shown[kind] == document_tools._viewer_ranges(ranges)
    read = _invoke(_tool(tools, "read_document"), handle=out["handle"],
                   pages="0-1")
    assert read["pages_returned"] == "0-1"
    assert read["pdf_pages_returned"] == "1-2"


def test_ranges_convert_and_cursors_are_left_alone():
    data = {"pages": "0-2,5", "n_pages": 9, "pages_with_hits": {"3": 2},
            "next": {"pages": "4-8", "start_line": 3},
            "segments": [{"id": 0, "pages": "0-1", "title": "x"}]}
    out = json.loads(document_tools.with_viewer_pages(json.dumps(data)))
    assert out["pdf_pages"] == "1-3,6"
    assert out["pdf_pages_with_hits"] == {"4": 2}
    assert "pdf_n_pages" not in out
    assert "pdf_pages" not in out["next"]              # a cursor: unchanged
    assert out["segments"][0]["pdf_pages"] == "1-2"


# ---------------------------------------------------------------------------
# C4: the marked copy
# ---------------------------------------------------------------------------

def _ready():
    if not document_tools.has_tool("annotate_document"):
        pytest.skip("installed planlens predates annotate_document")


def _two_marks(gt):
    return [
        {"kind": "highlight", "page": gt.narrative_page,
         "comment": "[#1] State the datum for these depths.",
         "quote": "20 to 35 feet"},
        {"kind": "box", "page": gt.narrative_page,
         "comment": "[#2] Which boring is this?", "bbox": [72, 300, 300, 330],
         "label": "#2"},
    ]


def test_rewriting_a_copy_this_conversation_holds_open(gt, folder):
    """F31 t4: the model had opened review_set_marked.pdf, then rebuilt it
    with append=false onto the same name; MuPDF could not remove the open
    file and the turn ended. Now the copy is swapped in, and the handle the
    model holds reads the NEW marks."""
    _ready()
    tools = make_vision_tools(engine=None, attachments={"s.pdf": gt.pdf})
    src = _invoke(_tool(tools, "open_document"), source="s.pdf")["handle"]
    annotate = _tool(tools, "annotate_document")
    first = _invoke(annotate, handle=src, output_path="s_marked.pdf",
                    markups=_two_marks(gt), check=False)
    assert first["n_written"] == 2 and first["output_path"] == "s_marked.pdf"
    held = _invoke(_tool(tools, "open_document"), source="s_marked.pdf")
    assert held["markups"]["n"] == 5 + 2 + 1          # theirs, ours, a label
    again = _invoke(annotate, handle=src, output_path="s_marked.pdf",
                    append=False, check=False, markups=[
                        {"kind": "note", "page": gt.narrative_page,
                         "comment": "[#1] HOLD POINT, not a suggestion.",
                         "point": [90, 90]}])
    assert "error" not in again, again
    assert again["output_path"] == "s_marked.pdf"
    marks = _invoke(_tool(tools, "document_markups"),
                    handle=held["handle"], author="AI draft")
    assert marks["n_markups"] == 1
    assert "HOLD POINT" in " ".join(marks["markups"])
    # no temporary file left behind, in the folder or its scratch
    leftovers = [n for root, _d, names in os.walk(folder) for n in names
                 if ".part" in n]
    assert leftovers == []


def test_a_second_call_adds_to_the_same_copy(gt, folder):
    _ready()
    tools = make_vision_tools(engine=None, attachments={"s.pdf": gt.pdf})
    src = _invoke(_tool(tools, "open_document"), source="s.pdf")["handle"]
    annotate = _tool(tools, "annotate_document")
    _invoke(annotate, handle=src, output_path="s_marked.pdf",
            markups=_two_marks(gt)[:1], check=False)
    held = _invoke(_tool(tools, "open_document"), source="s_marked.pdf")
    more = _invoke(annotate, handle=src, output_path="s_marked.pdf",
                   markups=_two_marks(gt)[1:], check=False)
    assert more["appended_to_existing"] is True
    ours = _invoke(_tool(tools, "document_markups"), handle=held["handle"],
                   author="AI draft")
    assert ours["n_markups"] == 3                    # highlight, box, label


def test_one_mark_is_replaced_by_its_id_and_the_rest_stay(gt, folder):
    _ready()
    tools = make_vision_tools(engine=None, attachments={"s.pdf": gt.pdf})
    src = _invoke(_tool(tools, "open_document"), source="s.pdf")["handle"]
    annotate = _tool(tools, "annotate_document")
    _invoke(annotate, handle=src, output_path="s_marked.pdf",
            markups=_two_marks(gt), check=False)
    copy = _invoke(_tool(tools, "open_document"), source="s_marked.pdf")
    listed = _invoke(_tool(tools, "document_markups"), handle=copy["handle"],
                     author="AI draft")["markups"]
    box_id = next(row.split("[", 1)[1].split("]", 1)[0] for row in listed
                  if "Which boring" in row)
    out = _invoke(annotate, handle=src, output_path="s_marked.pdf",
                  remove=[box_id], check=False, markups=[
                      {"kind": "box", "page": gt.narrative_page,
                       "comment": "[#2] HOLD POINT: name the boring.",
                       "bbox": [72, 300, 300, 330], "label": "#2 HOLD"}])
    assert [r["id"] for r in out["removed"]] == [box_id]
    assert out["removed"][0]["with_attached"] == 1         # its label
    assert out["n_written"] == 1
    after = _invoke(_tool(tools, "document_markups"), handle=copy["handle"],
                    author="AI draft")["markups"]
    text = " ".join(after)
    assert "Which boring" not in text and "name the boring" in text
    assert "State the datum" in text                       # the other stays
    # the reviewer's own five markups are still in the copy
    everyone = _invoke(_tool(tools, "document_markups"),
                       handle=copy["handle"])
    assert everyone["n_markups"] == 5 + 2 + 1


def test_remove_by_words_and_what_is_refused(gt, folder):
    _ready()
    tools = make_vision_tools(engine=None, attachments={"s.pdf": gt.pdf})
    src = _invoke(_tool(tools, "open_document"), source="s.pdf")["handle"]
    annotate = _tool(tools, "annotate_document")
    _invoke(annotate, handle=src, output_path="s_marked.pdf",
            markups=_two_marks(gt), check=False)
    copy = _invoke(_tool(tools, "open_document"), source="s_marked.pdf")
    theirs = _invoke(_tool(tools, "document_markups"),
                     handle=copy["handle"])["markups"]
    reviewer_id = next(row.split("[", 1)[1].split("]", 1)[0] for row in theirs
                       if gt.reviewer in row)
    out = _invoke(annotate, handle=src, output_path="s_marked.pdf",
                  check=False, remove=["state the datum", reviewer_id, "[#"])
    assert [r["says"][:4] for r in out["removed"]] == ["[#1]"]
    reasons = " ".join(r["reason"] for r in out["not_removed"])
    assert "not one this app wrote" in reasons          # the reviewer's
    assert "give the id" in reasons or "no mark" in reasons
    assert "ids" in out["ids_note"]


def test_removing_from_a_copy_that_does_not_exist_is_refused(gt, folder):
    _ready()
    tools = make_vision_tools(engine=None, attachments={"s.pdf": gt.pdf})
    src = _invoke(_tool(tools, "open_document"), source="s.pdf")["handle"]
    out = _invoke(_tool(tools, "annotate_document"), handle=src,
                  output_path="never.pdf", remove=["p0.m0"])
    assert "does not exist" in out["error"]


def test_rebuilding_the_copy_from_itself_is_refused(gt, folder):
    _ready()
    tools = make_vision_tools(engine=None, attachments={"s.pdf": gt.pdf})
    src = _invoke(_tool(tools, "open_document"), source="s.pdf")["handle"]
    annotate = _tool(tools, "annotate_document")
    _invoke(annotate, handle=src, output_path="s_marked.pdf",
            markups=_two_marks(gt), check=False)
    copy = _invoke(_tool(tools, "open_document"),
                   source="s_marked.pdf")["handle"]
    out = _invoke(annotate, handle=copy, output_path="s_marked.pdf",
                  append=False, check=False, markups=_two_marks(gt))
    assert "marked copy itself" in out["error"]
    assert "remove" in out["hint"]


def test_a_copy_open_elsewhere_is_saved_under_a_new_name(gt, folder,
                                                         monkeypatch):
    """A file the user has open in a viewer on a Windows host cannot be
    replaced: the copy is saved beside it and the result says so."""
    _ready()
    real_replace = os.replace
    calls = {"n": 0}

    def locked(src, dst):
        if str(dst).endswith("s_marked.pdf"):
            calls["n"] += 1
            raise PermissionError(13, "in use")
        return real_replace(src, dst)

    tools = make_vision_tools(engine=None, attachments={"s.pdf": gt.pdf})
    src = _invoke(_tool(tools, "open_document"), source="s.pdf")["handle"]
    monkeypatch.setattr(os, "replace", locked)
    monkeypatch.setattr(document_tools.time if hasattr(document_tools, "time")
                        else __import__("time"), "sleep", lambda s: None)
    out = _invoke(_tool(tools, "annotate_document"), handle=src,
                  output_path="s_marked.pdf", markups=_two_marks(gt),
                  check=False)
    assert out["output_path"] == "s_marked_2.pdf"
    assert "in use elsewhere" in out["saved_elsewhere"]
    assert (folder / "s_marked_2.pdf").is_file()
    assert calls["n"] == 3
