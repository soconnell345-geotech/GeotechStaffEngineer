"""Live smoke wave 1 fixes on the deep agent's tool surface.

* A12 -- ``plot_data`` is a direct tool on both pages, and ``write_xlsx``
  writes real Excel workbooks (rows or Markdown tables, one sheet each).
* A15a -- the scratch-space note fires on a ``glob`` that returns ``[]``.
* A15g -- a module method called as if it were a tool is answered with the
  ``call_agent`` form.
"""

import json
import os

import pytest

from funhouse_agent import _fileio, xlsx_writer
from funhouse_agent.deep.tools import make_vision_tools

pytestmark = pytest.mark.skipif(not xlsx_writer.available(),
                                reason="openpyxl not installed")


@pytest.fixture
def work(tmp_path, monkeypatch):
    monkeypatch.delenv(_fileio.DEFAULT_OUTPUT_DIR_ENV, raising=False)
    folder = tmp_path / "files"
    folder.mkdir()
    with _fileio.working_dir_bound(str(folder)):
        yield folder


def _tools(**kw):
    return {t.name: t for t in make_vision_tools(**kw)}


# ---------------------------------------------------------------------------
# write_xlsx
# ---------------------------------------------------------------------------

MD = """# Quantities

| Item | Qty | Unit |
|---|---:|---|
| Rebar #5 | 1,250 | ft |
| Concrete | **42.5** | cy |
| Tag | 007 | ea |

## Comment log

| No. | Sheet | Comment |
|:--|:--|:--|
| 1 | S-101 | Missing \\| escaped pipe |
"""


def test_markdown_tables_become_sheets_with_numbers():
    from openpyxl import load_workbook
    import io
    data, summary, warnings = xlsx_writer.build_workbook(markdown=MD)
    assert [s["sheet"] for s in summary] == ["Quantities", "Comment log"]
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Quantities"]
    assert [c.value for c in ws[1]] == ["Item", "Qty", "Unit"]
    assert ws["B2"].value == 1250 and ws["B3"].value == 42.5
    assert ws["B4"].value == "007"                 # an id stays text
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"
    assert wb["Comment log"]["C2"].value == "Missing | escaped pipe"
    assert not warnings


def test_rows_and_objects_and_sheet_names():
    from openpyxl import load_workbook
    import io
    data, summary, _ = xlsx_writer.build_workbook(sheets=[
        {"name": "Borings: B/1*", "rows": [["Depth (m)", "N"], [1.5, 12],
                                           ["3.0", "15"]]},
        {"name": "Borings: B/1*", "rows": [{"id": "B-1", "depth": 10},
                                           {"id": "B-2", "gwl": 3.2}]}])
    names = [s["sheet"] for s in summary]
    assert names[0] == "Borings- B-1-" and names[1] == "Borings- B-1- (2)"
    wb = load_workbook(io.BytesIO(data))
    assert wb[names[0]]["A3"].value == 3.0
    assert [c.value for c in wb[names[1]][1]] == ["id", "depth", "gwl"]
    assert wb[names[1]]["C3"].value == 3.2


def test_no_table_is_an_error():
    with pytest.raises(ValueError, match="no table to write"):
        xlsx_writer.build_workbook(markdown="just prose, no table")


def test_the_tool_saves_through_the_host_and_names_the_file(work):
    from webapp import core
    artifacts = []
    save_fn = core.make_save_fn(str(work), artifacts)
    tool = _tools(save_fn=save_fn)["write_xlsx"]
    assert "Excel" in tool.description and "markdown" in tool.description
    out = json.loads(tool.invoke({"path": "quantities.xls", "markdown": MD}))
    assert out["saved"] == "quantities.xlsx"
    assert (work / "quantities.xlsx").is_file()
    assert artifacts == [str(work / "quantities.xlsx")]
    assert "ATTACHED" in out["note"]
    assert str(work) not in json.dumps(out)
    bad = json.loads(tool.invoke({"path": "x", "markdown": "no table"}))
    assert "no table to write" in bad["error"]


# ---------------------------------------------------------------------------
# plot_data, directly
# ---------------------------------------------------------------------------

def test_plot_data_is_a_direct_tool_on_both_pages(work):
    from funhouse_agent.deep.agent import build_primary_tools
    from funhouse_agent.deep.review_agent import READER_TOOLS, REVIEW_TOOLS
    review = {t.name for t in build_primary_tools(allowed_agents=())}
    geotech = {t.name for t in build_primary_tools()}
    for names in (review, geotech):
        assert {"plot_data", "write_xlsx"} <= names
    assert "call_agent" not in review
    assert {"plot_data", "write_xlsx"} <= set(REVIEW_TOOLS)
    assert not ({"plot_data", "write_xlsx"} & set(READER_TOOLS))


def test_plot_data_describes_itself_and_draws(work):
    tool = _tools()["plot_data"]
    desc = tool.description
    assert "series" in desc and "INTERACTIVE" in desc and "SVG" in desc
    out = json.loads(tool.invoke({
        "series": [{"x": [0, 5, 10], "y": [0, 2, 3], "label": "N"}],
        "xlabel": "N (blows/ft)", "ylabel": "Depth (m)",
        "depth_axis": True, "output_path": "spt.png"}))
    assert "error" not in out, out
    assert out["output_path"] == "spt.png"
    assert (work / "spt.png").is_file()
    assert (work / "spt.plotly.json").is_file()
    assert str(work) not in json.dumps(out)


def test_read_tools_describe_the_conversation_not_the_server_disk():
    """Reads are confined to the conversation (plus the reference documents
    and any folder the deployment opens): no description advertises
    browsing /Workspace, /Volumes or /tmp, and the scratch-image names
    analyze_image takes are named."""
    tools = _tools()
    for name in ("list_files", "read_pdf_text", "read_text_file",
                 "analyze_image"):
        desc = tools[name].description
        for path in ("/Workspace", "/Volumes", "/tmp"):
            assert path not in desc, (name, path)
    assert "working folder" in tools["list_files"].description
    assert ".scratch" in tools["analyze_image"].description


# ---------------------------------------------------------------------------
# A15a: an empty glob
# ---------------------------------------------------------------------------

def test_an_empty_list_answer_gets_the_scratch_note():
    from langchain_core.messages import ToolMessage
    from funhouse_agent.deep.scratch_guard import (EMPTY_SCRATCH_NOTE,
                                                   ScratchFilesystemGuard)

    class Req:
        tool_call = {"name": "glob", "args": {"pattern": "*.pdf"}, "id": "1"}
        state = {}

    guard = ScratchFilesystemGuard()
    for empty in ("[]", "[ ]", []):
        out = guard.wrap_tool_call(Req(), lambda r, e=empty: ToolMessage(
            content=e, tool_call_id="1", name="glob"))
        text = out.content if isinstance(out.content, str) else json.dumps(
            out.content)
        assert EMPTY_SCRATCH_NOTE[:40] in text
    found = guard.wrap_tool_call(Req(), lambda r: ToolMessage(
        content="['/notes.md']", tool_call_id="1", name="glob"))
    assert found.content == "['/notes.md']"


# ---------------------------------------------------------------------------
# A15g: a module method called as a tool
# ---------------------------------------------------------------------------

def test_method_hint_names_the_call_agent_form():
    from funhouse_agent.deep.unknown_tool import method_hint
    from funhouse_agent.dispatch import ANALYSIS_MODULES
    hint = method_hint("write_diggs", ANALYSIS_MODULES)
    assert "call_agent(agent_name='subsurface', method='write_diggs'" in hint
    assert "describe_method('subsurface', 'write_diggs')" in hint
    assert method_hint("no_such_thing_xyz", ANALYSIS_MODULES) == ""
    # a module outside the scope is never suggested
    assert method_hint("write_diggs", ("bearing_capacity",)) == ""


def test_the_unknown_tool_error_gains_the_hint_in_a_real_agent(work):
    from langchain_core.language_models.chat_models import BaseChatModel
    from langchain_core.messages import AIMessage
    from langchain_core.outputs import ChatGeneration, ChatResult
    from funhouse_agent.deep.agent import build_deep_agent

    replies = [("", [("write_diggs", {"investigations": []})]),
               ("done", [])]

    class Scripted(BaseChatModel):
        @property
        def _llm_type(self):
            return "scripted"

        def bind_tools(self, tools, **kw):
            return self

        def _generate(self, messages, stop=None, run_manager=None, **kw):
            text, calls = replies.pop(0)
            return ChatResult(generations=[ChatGeneration(message=AIMessage(
                content=text, tool_calls=[{"name": n, "args": a, "id": "x1"}
                                          for n, a in calls]))])

    agent = build_deep_agent(Scripted())
    out = agent.invoke({"messages": [{"role": "user", "content": "q"}]})
    tool_msgs = [m for m in out["messages"] if m.type == "tool"]
    assert tool_msgs, out["messages"]
    text = tool_msgs[0].content
    assert "is not a valid tool" in text
    assert "call_agent(agent_name='subsurface', method='write_diggs'" in text
