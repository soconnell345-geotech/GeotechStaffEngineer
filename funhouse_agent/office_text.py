"""Word and Excel files as Markdown — the one reader every tool shares.

Live smoke wave 2a (B3): nothing read a .docx or an .xlsx. ``read_text_file``
refused them as binary; ``open_document`` rendered a .docx to pages and read
it back with tables one cell per line and list markers and line breaks gone,
so every "edit the memo" rebuilt the memo from flat text and drifted (F11 v2
lost its numbered list; F02's labels ran together); F31 invented a
``read_docx_placeholder`` tool; F41 met an .xlsx nothing could read.

* A **.docx** comes back in the dialect ``write_docx`` writes
  (:func:`calc_package.docx_renderer.docx_to_markdown`, the writer's inverse):
  headings, numbered and bulleted lists with their levels, pipe tables and
  their captions, bold / italic / code, line breaks, ``>`` quotes, fenced
  code, ``---`` page breaks and pictures — so read -> edit -> write keeps the
  document's shape.
* An **.xlsx / .xlsm** comes back as one ``## Sheet: <name>`` pipe table per
  sheet with a ``#`` column of Excel row numbers, values as Excel last
  calculated them (a formula never calculated shows as its formula). Its
  header is the first row filled across the table, mostly with text
  (:func:`_header_index`); rows above it -- a title, a project line -- are
  listed above the table, and the text says which row is the header (live
  smoke 2c, E7: F41's log has its title in row 1 and its header in row 4,
  and the reader said "row 1 is the header").

:func:`office_to_markdown` is the plain text; :func:`read_office` adds the
fields a tool result carries (format, title, sheets, pictures, a note). Both
raise :class:`OfficeReadError`, with a message fit to show, for a file they
cannot read. Library use only: no model, no network, nothing on disk but the
pictures a caller asks to be saved.
"""

from __future__ import annotations

import os
import re
from typing import Any, Dict, Optional, Tuple

#: What this module reads.
OFFICE_EXTENSIONS = (".docx", ".xlsx", ".xlsm")

#: Cells read from one workbook at most (a bigger one says where it stopped).
XLSX_MAX_CELLS = 100_000

#: The older binary Office formats: what to ask for instead.
LEGACY_OFFICE = {
    ".doc": "this is the old binary .doc format, which nothing here reads; "
            "ask for the file saved as .docx (Word: File > Save As > Word "
            "Document)",
    ".xls": "this is the old binary .xls format, which nothing here reads; "
            "ask for the file saved as .xlsx (Excel: File > Save As > Excel "
            "Workbook)",
}


class OfficeReadError(ValueError):
    """A Word or Excel file that could not be read; the message says why in
    words fit to show the user."""


def is_office_file(path) -> bool:
    """Whether ``path`` names a file this module reads (by extension)."""
    return os.path.splitext(str(path or ""))[1].lower() in OFFICE_EXTENSIONS


def office_to_markdown(path, *, image_dir: Optional[str] = None,
                       image_ref: Optional[str] = None) -> str:
    """The Markdown of a .docx / .xlsx / .xlsm (see the module docstring).

    ``image_dir``: a folder a .docx's pictures are saved into, so the
    Markdown references them (as ``<image_ref>/<stem>_figN.<ext>``,
    ``image_ref`` defaulting to ``image_dir``) and ``write_docx`` embeds them
    again; without it pictures are named in the text, not extracted.
    Raises :class:`OfficeReadError`."""
    return read_office(path, image_dir=image_dir, image_ref=image_ref)[0]


def read_office(path, *, image_dir: Optional[str] = None,
                image_ref: Optional[str] = None) -> Tuple[str, Dict[str, Any]]:
    """``(markdown, fields)``: :func:`office_to_markdown`'s text and what a
    tool result says about it — ``format``, ``read_as``, a ``note`` on how
    to use it, and ``title`` / ``pictures`` (docx) or ``sheets`` (xlsx).
    Raises :class:`OfficeReadError`."""
    path = str(path)
    ext = os.path.splitext(path)[1].lower()
    if ext in LEGACY_OFFICE:
        raise OfficeReadError(f"'{os.path.basename(path)}': "
                              f"{LEGACY_OFFICE[ext]}.")
    if ext not in OFFICE_EXTENSIONS:
        raise OfficeReadError(
            f"'{os.path.basename(path)}' is not a Word (.docx) or Excel "
            f"(.xlsx, .xlsm) file.")
    if not os.path.isfile(path):
        raise OfficeReadError(f"'{os.path.basename(path)}' does not exist.")
    kind = "a Word document" if ext == ".docx" else "an Excel workbook"
    try:
        if ext == ".docx":
            return _docx(path, image_dir, image_ref)
        return _xlsx(path)
    except ImportError as exc:
        raise OfficeReadError(
            f"'{os.path.basename(path)}' cannot be read here: {exc}") from exc
    except OfficeReadError:
        raise
    except Exception as exc:  # noqa: BLE001 - a damaged or odd file
        raise OfficeReadError(
            f"'{os.path.basename(path)}' could not be read as {kind}: "
            f"{type(exc).__name__}: {exc}") from exc


# -- Word ---------------------------------------------------------------------

def _docx(path, image_dir, image_ref):
    from calc_package.docx_renderer import docx_to_markdown
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_",
                  os.path.splitext(os.path.basename(path))[0]).strip("_")
    read = docx_to_markdown(path, image_dir=image_dir,
                            image_ref=image_ref if image_dir else None,
                            image_prefix=f"{stem or 'docx'}_")
    fields: Dict[str, Any] = {"format": "docx", "read_as": "markdown"}
    if read.get("title"):
        fields["title"] = read["title"]
    pictures = read.get("images") or []
    saved = [p for p in pictures if not p.startswith("picture ")]
    note = ("A Word document read as Markdown in write_docx's own dialect: "
            "headings, numbered and bulleted lists, pipe tables (an italic "
            "line under one is its caption), bold/italic/code, line breaks, "
            "> quotes, ``` code, and --- for a page break. To revise it, "
            "edit this Markdown and write it back with write_docx")
    if read.get("title"):
        note += (f" (title='{read['title']}': its Title is the first # "
                 f"line, written once)")
    note += "; the layout carries over."
    if saved:
        fields["pictures"] = saved
        note += (" Its pictures were saved beside it and are referenced by "
                 "name, so write_docx puts them back.")
    if len(saved) < len(pictures):
        fields["pictures_not_extracted"] = len(pictures) - len(saved)
        note += (" Pictures marked 'not extracted' are only named: they "
                 "will not be in a rewritten file.")
    fields["note"] = note
    return read["markdown"], fields


# -- Excel --------------------------------------------------------------------

def _cell_text(v) -> str:
    """One spreadsheet value as table-cell text."""
    import datetime as _dt
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 \
            else f"{v:.12g}"
    if isinstance(v, _dt.datetime):
        return (v.date().isoformat() if v.time() == _dt.time(0)
                else v.isoformat(sep=" ", timespec="minutes"))
    if isinstance(v, (_dt.date, _dt.time)):
        return v.isoformat()
    s = str(v).replace("\r\n", "\n").strip()
    return s.replace("|", r"\|").replace("\n", "<br>")


_NUMBER_TEXT = re.compile(r"^[+-]?[\d,]*\.?\d+(?:[eE][+-]?\d+)?%?$")


def _is_words(v) -> bool:
    """Whether a spreadsheet value is words -- a header's kind of cell --
    rather than a number, a date, a flag or a formula (a number stored as
    text counts as a number)."""
    if not isinstance(v, str):
        return False
    s = v.strip()
    return bool(s) and not s.startswith("=") \
        and not _NUMBER_TEXT.match(s.replace(" ", ""))


def _header_index(grid) -> Tuple[int, bool]:
    """``(index, found)`` of a sheet's header among its non-empty rows
    ``grid`` (``(row number, cells, words)`` each, ``words`` saying which
    cells are words): the first row filled across the table -- at least
    three fifths of the widest row's cells, and never fewer than two when
    the widest has two or more -- if most of its cells are words. Rows less
    filled above it are a title or a project line. When the first row
    filled across holds mostly numbers there is no header row, and the
    first row is taken (``found`` false)."""
    fills = [sum(1 for v in cells if v) for _n, cells, _w in grid]
    widest = max(fills)
    need = widest if widest <= 2 else max(2, -(-widest * 3 // 5))
    for i, (_n, cells, words) in enumerate(grid):
        filled = [k for k, v in enumerate(cells) if v]
        if len(filled) < need:
            continue
        if 2 * sum(1 for k in filled if words[k]) > len(filled):
            return i, True
        break
    return 0, False


def _xlsx(path):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    values = load_workbook(path, read_only=True, data_only=True)
    try:
        formulas = load_workbook(path, read_only=True, data_only=False)
    except Exception:  # noqa: BLE001 - values alone still read
        formulas = None
    parts, sheets = [], []
    budget = XLSX_MAX_CELLS
    stopped_at = None
    try:
        for ws in values.worksheets:
            if budget <= 0:
                stopped_at = stopped_at or ws.title
                sheets.append({"name": ws.title, "read": False})
                continue
            wf = None
            if formulas is not None:
                try:
                    wf = formulas[ws.title]
                except KeyError:
                    wf = None
            f_rows = wf.iter_rows(values_only=True) if wf is not None else None
            grid, cut = [], False
            for r_no, row in enumerate(ws.iter_rows(values_only=True), 1):
                f_row = next(f_rows, None) if f_rows is not None else None
                cells, words = [], []
                for j, v in enumerate(row):
                    if v is None and f_row is not None and j < len(f_row) \
                            and isinstance(f_row[j], str) \
                            and f_row[j].startswith("="):
                        v = f_row[j]
                    cells.append(_cell_text(v))
                    words.append(_is_words(v))
                grid.append((r_no, cells, words))
                budget -= max(1, len(cells))
                if budget <= 0:
                    cut = True
                    break
            grid = [(n, c, w) for n, c, w in grid if any(c)]
            hidden = getattr(ws, "sheet_state", "visible") != "visible"
            title = f"## Sheet: {ws.title}" + (" (hidden)" if hidden else "")
            if not grid:
                parts.append(title + "\n\n(empty)")
                sheets.append({"name": ws.title, "rows": 0, "columns": 0,
                               **({"hidden": True} if hidden else {})})
                continue
            # The header is found, not assumed (E7); a title or project line
            # above it is kept as text above the table.
            head_i, found = _header_index(grid)
            above, table = grid[:head_i], grid[head_i:]
            used = [k for _n, c, _w in table for k, v in enumerate(c) if v]
            c0, c1 = min(used), max(used) + 1
            width = c1 - c0
            rows = [(n, (c + [""] * c1)[c0:c1]) for n, c, _w in table]
            head_no, head = rows[0]
            why = ("the first row filled across the table, mostly with text"
                   if found else "no row filled across the table is mostly "
                   "text, so the first row is taken")
            lines = [title, ""]
            for n, c, _w in above:
                lines.append(f"Row {n}, above the table: "
                             + "; ".join(v for v in c if v))
            if above:
                lines.append("")
            lines += [f"Excel rows {head_no}-{rows[-1][0]}, "
                      f"columns {get_column_letter(c0 + 1)}-"
                      f"{get_column_letter(c1)}; row {head_no} is the header "
                      f"({why}).",
                      "",
                      "| # | " + " | ".join(head) + " |",
                      "| --- | " + " | ".join(["---"] * width) + " |"]
            lines += [f"| {n} | " + " | ".join(c) + " |" for n, c in rows[1:]]
            if cut:
                lines.append("")
                lines.append(f"(the reader stopped after Excel row "
                             f"{rows[-1][0]}: the workbook is larger than "
                             f"{XLSX_MAX_CELLS:,} cells)")
                stopped_at = ws.title
            parts.append("\n".join(lines))
            sheets.append({"name": ws.title, "rows": len(rows),
                           "columns": width, "header_row": head_no,
                           **({"rows_above_header": len(above)} if above
                              else {}),
                           **({"hidden": True} if hidden else {})})
    finally:
        values.close()
        if formulas is not None:
            formulas.close()
    note = ("An Excel workbook read as Markdown: one '## Sheet:' table per "
            "sheet, values as Excel last calculated them (a formula never "
            "calculated shows as its formula), the '#' column the Excel row "
            "number. Each table's header is the first row filled across it, "
            "mostly with text (the line above the table says which row); "
            "rows above the header, such as a title, are listed before the "
            "table. write_xlsx writes a new workbook.")
    if stopped_at:
        note += (f" Reading stopped in sheet '{stopped_at}' at "
                 f"{XLSX_MAX_CELLS:,} cells; later sheets are listed, "
                 f"not read.")
    return "\n\n".join(parts) + "\n", {"format": "xlsx",
                                       "read_as": "markdown",
                                       "sheets": sheets, "note": note}


__all__ = ["OFFICE_EXTENSIONS", "LEGACY_OFFICE", "XLSX_MAX_CELLS",
           "OfficeReadError", "is_office_file", "office_to_markdown",
           "read_office"]
