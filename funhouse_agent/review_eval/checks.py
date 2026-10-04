"""Deterministic checks on a review answer and the files it produced.

Each check is a dict with a ``type`` and its parameters; :func:`run_check`
returns ``(passed, detail)``. A check with ``"info": true`` is reported but not
scored. No check calls a model, so a score means the same thing on every run.

Terms. Wherever a check takes ``terms``, each term is one of:

* a string — found if it appears in the answer (case-insensitive, after
  :func:`normalize`: curly quotes, primes and dashes made ASCII, runs of
  whitespace collapsed);
* a list of strings/regexes — ALTERNATIVES, found if any one is;
* ``{"re": "..."}`` — a regular expression searched in the normalized answer
  (case-insensitive).
"""

from __future__ import annotations

import functools
import os
import re
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

_TRANSLATE = str.maketrans({
    "‘": "'", "’": "'", "′": "'", "´": "'",
    "“": '"', "”": '"', "″": '"',
    "‐": "-", "‑": "-", "‒": "-", "–": "-", "—": "-",
    "−": "-", " ": " ", " ": " ",
})


def normalize(text: str) -> str:
    """Lower-case ASCII-ish text for matching: quotes/primes/dashes made
    ASCII, markdown emphasis and backticks dropped, LaTeX number separators
    made plain (``1{,}195`` -> ``1,195``, ``1\\,195`` -> ``1195``; the
    Foundry rc2 run, 2026-10-03, failed a correct answer written in LaTeX),
    whitespace collapsed."""
    t = str(text or "").translate(_TRANSLATE).lower()
    t = t.replace("**", "").replace("__", "").replace("`", "")
    t = t.replace("{,}", ",")
    t = re.sub(r"(?<=\d)\\[,;:! ](?=\d)", "", t)
    return re.sub(r"\s+", " ", t)


def _dehyphen(text: str) -> str:
    """A hyphen between two letters read as a space: "edge-of-pavement"
    is "edge of pavement" (the Foundry run, 2026-10-02, failed four arms on
    a correct hyphenated answer). Digits are left alone (M-278, 2'-6")."""
    return re.sub(r"(?<=[a-z])-(?=[a-z])", " ", text)


def _one(term: Any, text: str) -> bool:
    if isinstance(term, dict) and "re" in term:
        return re.search(term["re"], text, flags=re.IGNORECASE) is not None
    want = normalize(str(term))
    return want in text or _dehyphen(want) in _dehyphen(text)


def term_found(term: Any, text: str) -> bool:
    """Whether ``term`` (string, alternatives list or regex) is in ``text``
    (already normalized)."""
    if isinstance(term, (list, tuple)):
        return any(_one(t, text) for t in term)
    return _one(term, text)


def _label(term: Any) -> str:
    if isinstance(term, (list, tuple)):
        return " | ".join(_label(t) for t in term)
    if isinstance(term, dict) and "re" in term:
        return f"/{term['re']}/"
    return str(term)


# ---------------------------------------------------------------------------
# Answer checks
# ---------------------------------------------------------------------------

def check_contains_all(answer: str, terms: Sequence[Any], **_) -> Tuple[bool, str]:
    text = normalize(answer)
    missing = [_label(t) for t in terms if not term_found(t, text)]
    return (not missing,
            "all present" if not missing else f"missing: {missing}")


def check_contains_any(answer: str, terms: Sequence[Any], **_) -> Tuple[bool, str]:
    text = normalize(answer)
    hit = [_label(t) for t in terms if term_found(t, text)]
    return (bool(hit), f"found: {hit}" if hit else
            f"none of: {[_label(t) for t in terms]}")


def check_not_contains(answer: str, terms: Sequence[Any], **_) -> Tuple[bool, str]:
    text = normalize(answer)
    hit = [_label(t) for t in terms if term_found(t, text)]
    return (not hit, "clean" if not hit else f"contains: {hit}")


def _mentions(text: str, item: str, aliases: Dict[str, Sequence[str]]) -> bool:
    forms = list(aliases.get(item) or [item])
    for f in forms:
        pat = re.escape(normalize(f))
        # An id must stand alone: "20.00a" must not match inside "120.00a",
        # and "11.01" must not match "11.015".
        if re.search(rf"(?<![\w.]){pat}(?![\w]|\.\d)", text):
            return True
    return False


def check_set_match(answer: str, vocabulary: Sequence[str],
                    expected: Sequence[str],
                    aliases: Optional[Dict[str, Sequence[str]]] = None,
                    min_recall: float = 1.0, min_precision: float = 0.0,
                    **_) -> Tuple[bool, str]:
    """The answer's mentions of ``vocabulary`` items, scored against the
    ``expected`` subset: recall and precision must both reach their floor.

    Mentions are counted over the whole answer, so an answer that names a
    sheet to rule it out ("30.01 prints 3600 cu ft, not concrete") is charged
    a false positive; ``min_precision`` is set with that in mind.
    """
    text = normalize(answer)
    aliases = aliases or {}
    said = {v for v in vocabulary if _mentions(text, v, aliases)}
    exp = set(expected)
    tp = len(said & exp)
    recall = tp / len(exp) if exp else 1.0
    precision = tp / len(said) if said else (1.0 if not exp else 0.0)
    ok = recall >= min_recall - 1e-9 and precision >= min_precision - 1e-9
    return ok, (f"recall {recall:.2f} (>= {min_recall}), precision "
                f"{precision:.2f} (>= {min_precision}); missing "
                f"{sorted(exp - said)}; extra {sorted(said - exp)}")


#: Where a later id cuts the text that belongs to an earlier one, at most.
LABEL_WINDOW = 160
#: Where the clause that belongs to an id ends (for "is another item's
#: title said of this id?").
_CLAUSE_END = re.compile(r"[.;](?:\s|$)")
#: The separators of an enumeration ("chapters 1, 2, 6 and 26").
_ENUM_SEP = r"\s*(?:,\s*and\b|,\s*&|,|\band\b|&|/|\bor\b)\s*"
_ENUM_RANGE = re.compile(r"^(\d+)\s*(?:-|to|through|thru)\s*(\d+)$")
#: A numeric range in an enumeration is read as its members up to this span
#: ("chapters 11-13"); a wider one ("chapters 1-26") names only its ends.
ENUM_MAX_SPAN = 4


def _title_in(text: str, forms: Sequence[str]) -> bool:
    for f in forms:
        pat = re.escape(normalize(f))
        if re.search(rf"(?<![\w-]){pat}", text):
            return True
    return False


def _enumerated(text: str, lead: str, tokens: Dict[str, str]
                ) -> List[Tuple[int, int, str]]:
    """``(start, end, item)`` for each item named in an enumeration after
    ``lead`` ("chapters 1, 2, 6, 7 and 26"), ``tokens`` mapping each item to
    the token that names it there ("7")."""
    by_token = {normalize(str(t)): item for item, t in tokens.items()}
    token = r"\d+(?:\s*(?:-|to|through|thru)\s*\d+)?|[a-z]\b"
    out: List[Tuple[int, int, str]] = []
    for m in re.finditer(rf"{lead}((?:{token})(?:{_ENUM_SEP}(?:{token}))*)",
                         text):
        start = m.start(1)
        for t in re.finditer(token, m.group(1)):
            got = t.group(0)
            names = [got]
            rng = _ENUM_RANGE.match(got)
            if rng:
                a, b = int(rng.group(1)), int(rng.group(2))
                names = ([str(n) for n in range(a, b + 1)]
                         if 0 <= b - a <= ENUM_MAX_SPAN else [str(a), str(b)])
            for name in names:
                item = by_token.get(name)
                if item is not None:
                    out.append((start + t.start(), start + t.end(), item))
    return out


def check_labelled_set(answer: str, items: Dict[str, Dict[str, Any]],
                       expected: Optional[Sequence[str]] = None,
                       min_recall: float = 1.0, require_title: bool = True,
                       window: int = LABEL_WINDOW,
                       enum_lead: Optional[str] = None,
                       enum_tokens: Optional[Dict[str, str]] = None,
                       **_) -> Tuple[bool, str]:
    """Items named BY THEIR ID TOGETHER WITH WHAT THEY ARE: an appendix
    letter with its title, a table number with its title.

    ``items`` maps each item to ``{"ids": [regex, ...], "titles": [phrase,
    ...]}``. Every id occurrence in the (normalized) answer owns the text
    after it, up to the next id of any item or ``window`` characters; an item
    counts when one of its occurrences owns one of its own titles, or is
    written "<title> (<id>)". So a list whose letters are matched to the
    wrong topics, a summary of topics with no ids, and an answer that ids
    only some items all fall short. With ``require_title=False`` an id
    alone counts too - unless the clause it owns names ANOTHER item's title
    and not its own - and ``enum_lead`` / ``enum_tokens`` read a bare
    enumeration ("ASCE 7 chapters 1, 2, 6 and 26") as ids.
    """
    text = normalize(answer)
    names = list(items)
    occ: List[Tuple[int, int, str]] = []
    for item, spec in items.items():
        for pat in spec.get("ids") or []:
            for m in re.finditer(pat, text):
                occ.append((m.start(), m.end(), item))
    if enum_lead and enum_tokens:
        occ += _enumerated(text, enum_lead, enum_tokens)
    # Overlapping matches (one id read two ways) keep the first, longest.
    occ.sort(key=lambda o: (o[0], -(o[1] - o[0])))
    kept: List[Tuple[int, int, str]] = []
    for o in occ:
        if kept and o[0] < kept[-1][1]:
            continue
        kept.append(o)
    said, mismatched = set(), set()
    for i, (start, end, item) in enumerate(kept):
        nxt = kept[i + 1][0] if i + 1 < len(kept) else len(text)
        after = text[end:min(nxt, end + int(window))]
        prev = kept[i - 1][1] if i else 0
        before = text[max(prev, start - int(window)):start]
        own = items[item].get("titles") or []
        if _title_in(after, own) or any(
                re.search(rf"(?<![\w-]){re.escape(normalize(f))}"
                          rf"[^.;|()\[\]]{{0,40}}[(\[][^.;|()\[\]]{{0,15}}$",
                          before) for f in own):
            said.add(item)
            continue
        clause = _CLAUSE_END.split(after, maxsplit=1)[0]
        other = [n for n in names if n != item and _title_in(
            clause, items[n].get("titles") or [])]
        if other:
            mismatched.add(item)
        elif not require_title:
            said.add(item)
    exp = set(expected if expected is not None else names)
    tp = len(said & exp)
    recall = tp / len(exp) if exp else 1.0
    ok = recall >= min_recall - 1e-9
    return ok, (f"recall {recall:.2f} (>= {min_recall}); named with what "
                f"they are: {sorted(said & exp)}; missing "
                f"{sorted(exp - said)}"
                + (f"; named with another item's title: "
                   f"{sorted(mismatched - said)}" if mismatched - said
                   else ""))


#: "page"/"p."/"pp." as a word of its own ("step. 12" and "mpg 30" are not
#: page citations).
_PAGE_WORD = r"(?<![a-z])(?:pdf\s+)?(?:pages?|pgs?\.?|pp\.|p\.)"
#: What separates the numbers in "pages 3, 5, and 7" / "pp. 3-7".
_PAGE_SEP = r"(?:,\s*and|,|and|&|-|to)"


def check_cites(answer: str, pages: Sequence[Any] = (),
                printed: Sequence[str] = (), sheets: Sequence[str] = (),
                **_) -> Tuple[bool, str]:
    """The answer cites the right place: a viewer page number (``pages``,
    1-based, written after "page"/"p."), a printed page number (``printed``,
    e.g. "5-2", anywhere) or a sheet id (``sheets``, anywhere)."""
    text = normalize(answer)
    for p in pages:
        if re.search(rf"{_PAGE_WORD}\s*(?:\d+\s*{_PAGE_SEP}\s*)*{int(p)}\b",
                     text):
            return True, f"cites page {p}"
    for p in printed:
        if re.search(rf"(?<![\w.-]){re.escape(normalize(p))}(?![\w])", text):
            return True, f"cites printed page {p}"
    for s in sheets:
        if _mentions(text, s, {}):
            return True, f"cites sheet {s}"
    return False, (f"no citation of pages {list(pages)}, printed "
                   f"{list(printed)} or sheets {list(sheets)}")


# ---------------------------------------------------------------------------
# File checks (on the files the agent wrote into its working folder)
# ---------------------------------------------------------------------------

def _files(files: Iterable[str], ext: str = "",
           name_contains: str = "") -> List[str]:
    out = []
    for f in files or ():
        base = os.path.basename(f).lower()
        if ext and not base.endswith(ext.lower()):
            continue
        if name_contains and name_contains.lower() not in base:
            continue
        out.append(f)
    return out


def check_file_produced(answer: str, files: Sequence[str] = (), ext: str = "",
                        name_contains: str = "", **_) -> Tuple[bool, str]:
    hits = _files(files, ext, name_contains)
    return (bool(hits), f"produced: {[os.path.basename(h) for h in hits]}"
            if hits else f"no {ext or 'file'} "
            f"{'named *' + name_contains + '* ' if name_contains else ''}"
            f"among {[os.path.basename(f) for f in files or ()]}")


def check_pdf_markups(answer: str, files: Sequence[str] = (), min: int = 1,
                      pages: Sequence[int] = (), text_contains: str = "",
                      author_contains: str = "", **_) -> Tuple[bool, str]:
    """A produced PDF carries at least ``min`` markups (on ``pages``, 0-based,
    if given; saying ``text_contains`` and signed ``author_contains`` if
    given), read back with planlens' own markup reader."""
    pdfs = _files(files, ".pdf")
    if not pdfs:
        return False, "no PDF produced"
    try:
        from planlens.document import Document
    except ImportError as exc:                       # pragma: no cover
        return False, f"planlens unavailable: {exc}"
    best = 0
    for path in pdfs:
        try:
            doc = Document(filepath=path)
            try:
                marks = doc.markups()
            finally:
                doc.close()
        except Exception as exc:                      # a broken PDF is a fail
            return False, f"{os.path.basename(path)} unreadable: {exc}"
        n = 0
        for m in marks:
            if pages and m.page not in pages:
                continue
            said = normalize(" ".join(str(x or "") for x in (
                getattr(m, "text", ""), getattr(m, "appearance_text", ""))))
            if text_contains and normalize(text_contains) not in said:
                continue
            if author_contains and normalize(author_contains) not in normalize(
                    getattr(m, "author", "") or ""):
                continue
            n += 1
        best = max(best, n)
    return best >= min, f"{best} matching markup(s) (need {min})"


#: "page 4", "pages 4, 12 and 20", "pp. 4-6", "PDF page 12", "sheets 3 & 9",
#: "PDF pages: 4, 12, 19, and 20" (the colon form cost a correct answer its
#: list in the Foundry rc2 run, 2026-10-03).
_PAGE_LIST = re.compile(
    r"\b(?:pdf\s+)?(?:pages?|pp?\.|sheets?)\s*[:#]?\s*"
    r"(\d+(?:\s*(?:-|–|to|through)\s*\d+)?"
    r"(?:\s*(?:,\s*and|,|and|&)\s*\d+(?:\s*(?:-|–|to|through)\s*\d+)?)*)")


def pages_named(answer: str, n_pages: int, max_span: int = 4) -> set:
    """The viewer page numbers (1..``n_pages``) an answer names after
    "page(s)", "p./pp." or "sheet(s)", enumerations and short ranges read."""
    found = set()
    for m in _PAGE_LIST.finditer(normalize(answer)):
        for part in re.split(r"\s*(?:,\s*and|,|and|&)\s*", m.group(1)):
            rng = re.match(r"(\d+)\s*(?:-|–|to|through)\s*(\d+)$", part.strip())
            if rng:
                a, b = int(rng.group(1)), int(rng.group(2))
                if 0 <= b - a <= max_span:
                    found.update(range(a, b + 1))
                continue
            if part.strip().isdigit():
                found.add(int(part))
    return {p for p in found if 1 <= p <= n_pages}


def check_pages_listed(answer: str, expected: Sequence[int] = (),
                       n_pages: int = 0, min_recall: float = 1.0,
                       min_precision: float = 0.75, **_) -> Tuple[bool, str]:
    """The pages the answer names (viewer numbering) against the pages that
    really carry the thing: a "which sheets have X" question across a long
    set, answerable only by looking at every sheet."""
    said = pages_named(answer, n_pages)
    exp = {int(p) for p in expected}
    tp = len(said & exp)
    recall = tp / len(exp) if exp else 1.0
    precision = tp / len(said) if said else 0.0
    ok = recall >= min_recall - 1e-9 and precision >= min_precision - 1e-9
    return ok, (f"pages named {sorted(said)}; expected {sorted(exp)}: recall "
                f"{recall:.2f} (>= {min_recall}), precision {precision:.2f} "
                f"(>= {min_precision})")


@functools.lru_cache(maxsize=4)
def _fixture_tags(name: str) -> tuple:
    """The ground-truth tags of a synthetic sheet set (built, not stored, so
    the truth cannot drift from the PDF the task is asked of)."""
    if name == "tags":
        from planlens.testing.tag_fixtures import build_synthetic_tag_set
        return tuple(build_synthetic_tag_set().tags)
    raise KeyError(f"no tag fixture {name!r}")


def _centre_inside(box: Sequence[float], mark: Sequence[float],
                   pad: float) -> bool:
    cx, cy = (box[0] + box[2]) / 2.0, (box[1] + box[3]) / 2.0
    return (mark[0] - pad <= cx <= mark[2] + pad
            and mark[1] - pad <= cy <= mark[3] + pad)


def _area(box: Sequence[float]) -> float:
    return max(0.0, box[2] - box[0]) * max(0.0, box[3] - box[1])


def check_markups_on_targets(answer: str, files: Sequence[str] = (),
                             fixture: str = "tags", text: str = "",
                             kind: str = "", page: Optional[int] = None,
                             mark_kinds: Sequence[str] = ("Circle", "Square"),
                             min_recall: float = 0.8,
                             min_precision: float = 0.8, pad: float = 4.0,
                             max_area_factor: float = 60.0,
                             **_) -> Tuple[bool, str]:
    """A produced PDF's rings and boxes sit ON the true targets.

    The targets are a fixture's ground-truth tags (``text`` / ``kind`` /
    ``page`` pick them). A mark hits a target when the target's centre is
    inside the mark (``pad`` points of slack) and the mark is not a blanket —
    at most ``max_area_factor`` times the target's area (25 pt² floor).
    Recall is the targets hit; precision the marks that hit any target. Field
    session 2026-10-01: circles drawn at invented coordinates passed every
    "is there a markup" check and were nowhere near the tags.
    """
    targets = [t for t in _fixture_tags(fixture)
               if (not text or t.text == text) and (not kind or t.kind == kind)
               and (page is None or t.page == page)]
    if not targets:
        return False, "the fixture has no such targets (a broken check)"
    pdfs = _files(files, ".pdf")
    if not pdfs:
        return False, "no PDF produced"
    from planlens.document import Document
    best = None
    for path in pdfs:
        try:
            doc = Document(filepath=path)
            try:
                marks = [m for m in doc.markups()
                         if m.kind in tuple(mark_kinds)
                         and (page is None or m.page == page)]
            finally:
                doc.close()
        except Exception as exc:
            return False, f"{os.path.basename(path)} unreadable: {exc}"

        def hit(m, t):
            return (m.page == t.page and _centre_inside(t.bbox, m.bbox, pad)
                    and _area(m.bbox) <= max_area_factor
                    * max(_area(t.bbox), 25.0))

        found = sum(1 for t in targets if any(hit(m, t) for m in marks))
        good = sum(1 for m in marks if any(hit(m, t) for t in targets))
        recall = found / len(targets)
        precision = good / len(marks) if marks else 0.0
        row = (recall, precision, found, good, len(marks),
               os.path.basename(path))
        if best is None or row[:2] > best[:2]:
            best = row
    recall, precision, found, good, n_marks, name = best
    ok = recall >= min_recall and precision >= min_precision
    return ok, (f"{name}: {found}/{len(targets)} targets marked (recall "
                f"{recall:.2f} >= {min_recall}), {good}/{n_marks} marks on a "
                f"target (precision {precision:.2f} >= {min_precision})")


def check_docx_contains(answer: str, files: Sequence[str] = (),
                        terms: Sequence[Any] = (), **_) -> Tuple[bool, str]:
    docs = _files(files, ".docx")
    if not docs:
        return False, "no .docx produced"
    try:
        import docx
    except ImportError as exc:                       # pragma: no cover
        return False, f"python-docx unavailable: {exc}"
    texts = []
    for path in docs:
        try:
            d = docx.Document(path)
        except Exception as exc:
            return False, f"{os.path.basename(path)} unreadable: {exc}"
        texts += [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                texts += [c.text for c in row.cells]
    ok, detail = check_contains_all("\n".join(texts), terms)
    return ok, detail


# ---------------------------------------------------------------------------
# Process checks (informational by default: outcomes are what is scored)
# ---------------------------------------------------------------------------

def check_tool_used(answer: str, tool_calls: Sequence[Dict[str, Any]] = (),
                    any: Sequence[str] = (), min_calls: int = 1,
                    **_) -> Tuple[bool, str]:
    names = [c.get("name") for c in tool_calls or ()]
    n = sum(1 for x in names if x in set(any))
    return n >= min_calls, f"{n} call(s) of {list(any)}"


CHECKS = {
    "contains_all": check_contains_all,
    "contains_any": check_contains_any,
    "not_contains": check_not_contains,
    "set_match": check_set_match,
    "labelled_set": check_labelled_set,
    "cites": check_cites,
    "file_produced": check_file_produced,
    "pdf_markups": check_pdf_markups,
    "markups_on_targets": check_markups_on_targets,
    "pages_listed": check_pages_listed,
    "docx_contains": check_docx_contains,
    "tool_used": check_tool_used,
}


def run_check(check: Dict[str, Any], answer: str,
              files: Sequence[str] = (),
              tool_calls: Sequence[Dict[str, Any]] = ()) -> Dict[str, Any]:
    """Run one check; never raises (a broken check is a failed check)."""
    kind = check.get("type")
    fn = CHECKS.get(kind)
    params = {k: v for k, v in check.items()
              if k not in ("type", "label", "info")}
    label = check.get("label") or kind
    if fn is None:
        return {"type": kind, "label": label, "passed": False,
                "info": bool(check.get("info")),
                "detail": f"unknown check type {kind!r}"}
    try:
        ok, detail = fn(answer or "", files=files, tool_calls=tool_calls,
                        **params)
    except Exception as exc:
        ok, detail = False, f"check error: {type(exc).__name__}: {exc}"
    return {"type": kind, "label": label, "passed": bool(ok),
            "info": bool(check.get("info")), "detail": detail}


def score(results: Sequence[Dict[str, Any]]) -> Dict[str, Any]:
    """Scored checks passed / total, and whether the task passed (all did)."""
    scored = [r for r in results if not r.get("info")]
    passed = sum(1 for r in scored if r["passed"])
    return {"checks_passed": passed, "checks_total": len(scored),
            "passed": bool(scored) and passed == len(scored)}


__all__ = ["normalize", "term_found", "run_check", "score", "CHECKS"]
