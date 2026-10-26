"""The owner's report question list as a schema, code's answers and the
answer sheet (BACKLOG RI-90). Offline: a made-up list and code map built in
tmp_path, a made-up record, a fake engine. Every question, key, value and
name here is invented; the owner's list is private and never read here.
"""

from __future__ import annotations

import json

import pytest

from report_ingest import question_list as ql
from report_ingest.engine import STRICT_UNSUPPORTED, strict_schema
from report_ingest.model import (BoundReport, Calculation, DocumentFacts,
                                 Investigation, LabTest, Layer,
                                 NamedQuantity, NarrativeFacts, PageLabel,
                                 Provenance, Quantity, ReportRecord,
                                 WaterLevel)
from report_ingest.question_list import Answer, Gate
from report_ingest.tests.fake_engine import FakeEngine

HEADINGS = ["Output", "Category", "Data Type", "Allow Null",
            "Other Validation", "Prompt", "Key", "Answered by",
            "Changed in v2"]

RULES = ["Rules for this list (a heading)",
         "Quote numbers with their units.",
         "Leave a row empty when unsure."]

#: A made-up list: one row of every Data Type, three screening rows, Code
#: rows and Both rows code can and cannot answer.
ROWS = [
    ("Soil study?", "Screening", "Boolean", None, None,
     "Is this a soil study?", "is_soil_study", "Both, compared"),
    ("Log book?", "Screening", "Boolean", None, None,
     "Is this only a book of logs?", "is_log_book", "Report text (model)"),
    ("Advice only?", "Screening", "Boolean", None, None,
     "Is this advice alone?", "is_advice_only", "Report text (model)"),
    ("Attached", "Screening", "List of records", True,
     "parts: title; kind [earlier volume, appended study, loose sheets]; "
     "pdf_pages", "List the attached studies.", "attached_studies", "Code"),
    ("Flat?", "Site ", "Boolean", True, None, "Is the site flat?",
     "site_flat", "Report text (model)"),
    ("Wells", "Site", "Integer", True, None, "How many wells?",
     "well_count", "Report text (model)"),
    ("Holes", "Site", "Integer", True, None, "How many holes were drilled?",
     "hole_total", "Both, compared"),
    ("Flood risk", "Hazards", "String", True, "[low, high, unknown]",
     "What flood risk?", "flood_risk", "Report text (model)"),
    ("Issued", "Paperwork", "Date", None, "[YYYY-MM]",
     "When was it issued?", "issue_month", "Report text (model)"),
    ("Features", "Hazards", "List of strings", True, "[creek, cliff]",
     "Which features are near?", "nearby_features", "Report text (model)"),
    ("Plan sheets", "Paperwork", "Integer or list of integers", True, None,
     "Which pages are plans?", "plan_sheet_pages", "Code"),
    ("Colours", "Site", "Dictionary", True,
     "parts: colour (value); name (key)", "Name the colours.",
     "colour_names", "Report text (model)"),
    ("Posts", "Advice", "List of records", True,
     "parts: post; height_text; finish [painted, bare]; pdf_page",
     "List the posts.", "post_heights", "Report text (model)"),
    ("Piles?", "Advice", "Boolean", None, None, "Are piles proposed?",
     "uses_piles", "Both, compared"),
    ("Page kinds", "Paperwork", "Dictionary", None,
     "parts: pdf_page; sheet_class; printed_page",
     "Give each page's kind.", "page_kinds", "Code"),
    ("Summary", "Site", "string", True, None, "Summarize the study.",
     "short_summary", "Report text (model)"),
    ("Water", "Site", "String", True, None, "Describe the water seen.",
     "water_seen", "Both, compared"),
    ("Back figures", "Paperwork", "Integer", True, None,
     "How many figures at the back?", "back_figure_total",
     "Both, compared"),
]

#: The made-up list's code map: is_soil_study and back_figure_total are
#: Both rows it gives no reading.
CODE_MAP = {
    "readings": {
        "attached_studies": "bound_documents",
        "hole_total": "hole_count:boring,well",
        "plan_sheet_pages": "pages_labelled:plan",
        "uses_piles": "calculation_of:axial_pile,lateral_pile,pile_group",
        "page_kinds": "page_labels",
        "water_seen": "water_readings",
    },
    "gates": [{"key": "is_soil_study", "closes_when": False, "closes": "*"},
              {"key": "is_log_book", "closes_when": True,
               "closes": "Advice"}],
    "not_both_yes": [["is_log_book", "is_advice_only"]],
}


def _workbook(path, rows=ROWS, rules=RULES, headings=HEADINGS,
              code_map=CODE_MAP):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Rules"
    for rule in rules:
        sheet.append([rule])
    questions = book.create_sheet("Questions")
    questions.append(headings)
    for row in rows:
        questions.append(list(row) + ([None] * (len(headings) - len(row))))
    book.save(path)
    if code_map is not None:
        path.with_name(path.stem + ql.CODE_MAP_SUFFIX).write_text(
            json.dumps(code_map), encoding="utf-8")
    return path


@pytest.fixture()
def qlist(tmp_path):
    return ql.load_questions(_workbook(tmp_path / "questions.xlsx"))


# ---------------------------------------------------------------------------
# loading the list and its code map
# ---------------------------------------------------------------------------

def test_the_list_loads_with_its_rules_and_code_map(qlist):
    assert len(qlist) == len(ROWS)
    assert qlist.rules == RULES[1:]               # the heading is not a rule
    assert qlist.name == "questions.xlsx" and qlist.modified is not None
    by = qlist.by_key()
    assert by["site_flat"].category == "Site"     # the trailing space goes
    assert by["short_summary"].data_type == "string"
    assert by["hole_total"].answered_by == "both"
    assert by["attached_studies"].answered_by == "code"
    assert by["site_flat"].answered_by == "model"
    assert by["flood_risk"].options == ("low", "high", "unknown")
    assert by["issue_month"].validation == "YYYY-MM"
    assert by["site_flat"].allow_null and not by["is_soil_study"].allow_null
    kind = by["attached_studies"].parts[1]
    assert (kind.name, len(kind.options)) == ("kind", 3)
    # A Dictionary's key part comes first; unmarked parts keep their order.
    assert [(p.name, p.role) for p in by["colour_names"].parts] == [
        ("name", "key"), ("colour", "value")]
    assert [p.name for p in by["page_kinds"].parts] == [
        "pdf_page", "sheet_class", "printed_page"]
    assert [q.key for q in ql.screening_questions(qlist)] == [
        "is_soil_study", "is_log_book", "is_advice_only", "attached_studies"]
    assert len(ql.main_questions(qlist)) == len(ROWS) - 4
    assert qlist.map_name == "questions.code.json"
    assert qlist.readings["hole_total"] == ("hole_count", "boring,well")
    assert qlist.gates == (Gate("is_soil_study", False, "*"),
                           Gate("is_log_book", True, "Advice"))
    assert qlist.not_both_yes == (("is_log_book", "is_advice_only"),)


def test_the_list_path_can_come_from_the_variable(tmp_path, monkeypatch):
    path = _workbook(tmp_path / "list.xlsx")
    monkeypatch.setenv(ql.LIST_ENV, str(path))
    assert len(ql.load_questions()) == len(ROWS)
    monkeypatch.delenv(ql.LIST_ENV)
    with pytest.raises(ql.QuestionListError, match=ql.LIST_ENV):
        ql.load_questions()


def test_without_a_code_map_code_reads_nothing_and_nothing_closes(tmp_path):
    bare = ql.load_questions(_workbook(tmp_path / "bare.xlsx",
                                       code_map=None))
    assert (bare.map_name, bare.readings, bare.gates) == ("", {}, ())
    code = ql.code_answers(_record(), bare)
    assert {a.note for a in code.values()} == {
        "not read: the list has no code map"}
    assert ql.closed_rows(bare, {"is_soil_study": Answer(
        value=False, status="answered")}) == {}


def test_every_malformed_row_is_named(tmp_path):
    bad = [
        ("A", "Site", "Float", None, None, "How deep?", "depth",
         "Report text (model)"),                                   # row 2
        ("B", "Site", "String", None, None, "Who?", "Who Wrote",
         "Report text (model)"),                                   # row 3
        ("C", "Site", "String", None, None, "Who?", "author",
         "Someone"),                                               # row 4
        ("D", "Site", "Integer", None, "[1, 2]", "How many?",
         "count", "Report text (model)"),                          # row 5
        ("E", "Site", "List of records", None, None, "List them.",
         "items", "Report text (model)"),                          # row 6
        ("F", "Site", "String", None, "parts: a; b", "What?", "what",
         "Report text (model)"),                                   # row 7
        ("G", "Site", "Boolean", "maybe", None, "Level?", "level",
         "Report text (model)"),                                   # row 8
        ("H", "Site", "Boolean", None, None, "Flat?", "flat",
         "Report text (model)"),                                   # row 9
        ("I", "Site", "Integer", None, None, "How many posts?", "posts",
         "Report text (model)"),                                   # row 10
        ("J", "Site", "Dictionary", None, "parts: a (key); b (key)",
         "Map them.", "mapping", "Report text (model)"),           # row 11
        ("K", "Site", "String", None, None, "", "empty",
         "Report text (model)"),                                   # row 12
        ("L", "Site", "Boolean", None, None, "Flat?", "flat",
         "Report text (model)"),                                   # row 13
    ]
    with pytest.raises(ql.QuestionListError) as caught:
        ql.load_questions(_workbook(tmp_path / "bad.xlsx", rows=bad,
                                    code_map=None))
    said = str(caught.value)
    for row, words in ((2, "Data Type 'Float'"), (3, "Key must be"),
                       (4, "Answered by 'Someone'"),
                       (5, "closed list, which only a String"),
                       (6, "needs its parts"), (7, "gives parts"),
                       (8, "Allow Null 'maybe'"),
                       (11, "one part '(key)' and one '(value)'"),
                       (12, "has no Prompt"),
                       (13, "already used on row 9")):
        line = next((ln for ln in said.splitlines()
                     if ln.startswith(f"row {row} ")), "")
        assert words in line, (row, said)
    assert not any(ln.startswith(("row 9 ", "row 10 "))
                   for ln in said.splitlines())


def test_a_code_map_that_does_not_fit_the_list_is_named(tmp_path):
    wrong = {
        "readings": {"no_such_key": "soft_layers",
                     "site_flat": "soft_layers",
                     "hole_total": "guess",
                     "water_seen": "hole_count:tunnel",
                     "plan_sheet_pages": "pages_labelled:nowhere",
                     "page_kinds": "page_labels:extra",
                     "uses_piles": "calculation_of:"},
        "gates": [{"key": "site_flat", "closes_when": True},
                  {"key": "is_log_book", "closes_when": "yes"},
                  {"key": "is_soil_study", "closes_when": False,
                   "closes": "Nowhere"}],
        "not_both_yes": [["is_log_book", "well_count"]],
    }
    with pytest.raises(ql.QuestionListError) as caught:
        ql.load_questions(_workbook(tmp_path / "x.xlsx", code_map=wrong))
    said = str(caught.value)
    for words in ("Key 'no_such_key': the list has no such Key",
                  "Key 'site_flat': the list gives that row to the model",
                  "Key 'hole_total': 'guess' is not one of",
                  "Key 'water_seen': give the kinds", "(not tunnel)",
                  "Key 'plan_sheet_pages': the page label 'nowhere'",
                  "Key 'page_kinds': this reading takes no argument",
                  "Key 'uses_piles': give the kinds",
                  "gate on Key 'site_flat': a gate is a Boolean row",
                  "gate on Key 'is_log_book': closes_when is true or false",
                  "no row is in the category 'Nowhere'",
                  "pair ['is_log_book', 'well_count']"):
        assert words in said, words
    (tmp_path / "x.code.json").write_text("{not json", encoding="utf-8")
    with pytest.raises(ql.QuestionListError, match="is not JSON"):
        ql.load_questions(tmp_path / "x.xlsx")
    with pytest.raises(ql.QuestionListError, match="code map is not a file"):
        ql.load_questions(tmp_path / "x.xlsx",
                          code_map=tmp_path / "missing.json")


def test_a_gate_or_pair_on_a_row_code_answers_is_refused(tmp_path):
    rows = ROWS + [("Volumes?", "Screening", "Boolean", None, None,
                    "Is it in volumes?", "has_volumes", "Code")]
    wrong = dict(CODE_MAP, gates=[{"key": "has_volumes",
                                   "closes_when": True, "closes": "*"}],
                 not_both_yes=[["is_log_book", "has_volumes"]])
    with pytest.raises(ql.QuestionListError) as caught:
        ql.load_questions(_workbook(tmp_path / "x.xlsx", rows=rows,
                                    code_map=wrong))
    said = str(caught.value)
    assert "gate on Key 'has_volumes': the list gives that row to code" \
        in said
    assert "pair ['is_log_book', 'has_volumes']: the list gives a row of " \
        "it to code" in said


def test_a_sheet_without_the_headings_is_refused(tmp_path):
    path = _workbook(tmp_path / "x.xlsx", headings=["Question", "Answer"],
                     rows=[("Is the site flat?", "")], code_map=None)
    with pytest.raises(ql.QuestionListError, match="no sheet has"):
        ql.load_questions(path)


# ---------------------------------------------------------------------------
# the answer schema
# ---------------------------------------------------------------------------

def _answer(value, status="answered", pages=(), note=""):
    return {"value": value, "status": status, "pdf_pages": list(pages),
            "note": note}


def _full(**changes):
    out = {
        "is_soil_study": _answer(True), "is_log_book": _answer(False),
        "is_advice_only": _answer(False),
        "attached_studies": _answer([]),
        "site_flat": _answer(None, "not said"),
        "well_count": _answer(2, pages=[4]), "hole_total": _answer(3),
        "flood_risk": _answer("low"), "issue_month": _answer("2020-01"),
        "nearby_features": _answer(["creek"]),
        "plan_sheet_pages": _answer([5]),
        "colour_names": _answer([{"name": "A", "colour": "Red"}]),
        "post_heights": _answer([{"post": "A", "height_text": "2 m",
                                  "finish": "bare", "pdf_page": 3}]),
        "uses_piles": _answer(None, "does not apply"),
        "page_kinds": _answer([{"pdf_page": 1, "sheet_class": "cover",
                                "printed_page": ""}]),
        "short_summary": _answer("A made-up summary."),
        "water_seen": _answer("Dry to the depths explored."),
        "back_figure_total": _answer(4),
    }
    out.update(changes)
    return out


def test_every_data_type_takes_its_shape(qlist):
    schema = ql.answer_schema(qlist)
    got = schema.model_validate(_full())
    assert got.colour_names.value[0].name == "A"
    assert got.post_heights.value[0].pdf_page == 3
    assert got.plan_sheet_pages.value == [5]
    assert got.site_flat.status == "not said"
    assert got.uses_piles.status == "does not apply"


@pytest.mark.parametrize("key, value", [
    ("flood_risk", "medium"),                       # off the closed list
    ("nearby_features", ["creek", "quarry"]),       # one item off the list
    ("post_heights", [{"post": "A", "height_text": "2 m",
                       "finish": "rusty", "pdf_page": 3}]),
    ("colour_names", {"A": "Red"}),                 # an open dictionary
    ("well_count", "two"),
    ("plan_sheet_pages", 5),                        # one page is a list
    ("site_flat", None),
])
def test_a_value_off_its_shape_is_refused(qlist, key, value):
    schema = ql.answer_schema(qlist)
    bad = _full(**{key: _answer(value)})
    if key == "site_flat":
        bad[key]["status"] = "maybe"                # not a status
    with pytest.raises(Exception):
        schema.model_validate(bad)


def test_a_part_s_type_follows_its_name(qlist):
    row = [qlist.by_key()["attached_studies"]]
    item = ql.answer_schema(row).model_validate({
        "attached_studies": _answer([{"title": "T", "kind": None,
                                      "pdf_pages": [4, 5]}])})
    assert item.attached_studies.value[0].pdf_pages == [4, 5]
    with pytest.raises(Exception):
        ql.answer_schema(row).model_validate({
            "attached_studies": _answer([{"title": "T", "kind": "a sequel",
                                          "pdf_pages": []}])})


def _nodes(schema):
    out = []

    def walk(node):
        if isinstance(node, list):
            for item in node:
                walk(item)
        elif isinstance(node, dict):
            if node.get("type") == "object" or "properties" in node:
                out.append(node)
            for value in node.values():
                walk(value)

    walk(schema)
    return out


def test_the_schema_is_acceptable_to_strict_structured_output(qlist):
    schema = strict_schema(ql.answer_schema(qlist))
    text = json.dumps(schema)
    assert not [k for k in STRICT_UNSUPPORTED if f'"{k}":' in text]
    objects = _nodes(schema)
    assert len(objects) > len(ROWS)
    for node in objects:
        # Closed, every field required, and no open dictionary: an object
        # with no properties would be one strict mode lets carry nothing.
        assert node["additionalProperties"] is False
        assert node["properties"]
        assert sorted(node["required"]) == sorted(node["properties"])
    # Every top-level answer is a plain $ref, and no node anywhere has a
    # $ref with another keyword beside it.
    assert all(set(p) == {"$ref"} for p in schema["properties"].values())
    beside = []

    def walk(node):
        if isinstance(node, list):
            for item in node:
                walk(item)
        elif isinstance(node, dict):
            if "$ref" in node and len(node) > 1:
                beside.append(node)
            for value in node.values():
                walk(value)

    walk(schema)
    assert beside == []
    # The question rides on its own answer model.
    ref = schema["properties"]["flood_risk"]["$ref"].rsplit("/", 1)[1]
    assert schema["$defs"][ref]["description"].startswith(
        "Flood risk: What flood risk?")


# ---------------------------------------------------------------------------
# a made-up record
# ---------------------------------------------------------------------------

def _q(value, unit="m", page=6):
    return Quantity(value=value, unit=unit,
                    prov=Provenance(page=page))


def _log(name, kind="boring", page=6, soft=False, water=()):
    return Investigation(
        investigation_id=name, kind=kind, depth_unit="m", pages=[page],
        layers=[Layer(top=_q(0.0, page=page), bottom=_q(2.0, page=page),
                      description=("very soft grey CLAY" if soft
                                   else "dense brown SAND"),
                      prov=Provenance(page=page))],
        water=[WaterLevel(depth=_q(d, page=page), when="while_drilling",
                          prov=Provenance(page=page)) for d in water])


def _record(**changes):
    labels = ["cover", "letter", "narrative", "narrative", "plan",
              "boring_log", "boring_log", "boring_log", "calculation",
              "lab_test"]
    record = ReportRecord(
        document=DocumentFacts(report_id="R00", n_pages=len(labels)),
        page_labels=[PageLabel(page=n, label=lab)
                     for n, lab in enumerate(labels)],
        narrative=NarrativeFacts(counted={"tables": 2, "figures": 0}),
        investigations=[_log("B-1", water=[1.5]),
                        _log("B-01", kind="well", page=6),
                        _log("B-2", page=7, soft=True),
                        _log("HA-1", kind="hand_auger", page=7)],
        calculations=[Calculation(kind="axial_pile", pages=[8],
                                  results=[NamedQuantity(name="Capacity",
                                                         text="OK")])],
        lab_tests=[LabTest(kind="chemical", investigation_id="B-1",
                           pages=[9],
                           result={"kind": "chemical", "pH": 7.5})],
        bound_documents=[
            BoundReport(bound_id="bound1", title="An earlier study",
                        firm="Firm A", date="2001",
                        kind="appended_prior_report", first_page=8,
                        last_page=8, n_pages=1),
            BoundReport(bound_id="bound2", title="Some logs",
                        kind="data_report", first_page=9, last_page=9,
                        n_pages=1)])
    return record.model_copy(update=changes)


def test_code_answers_from_the_record(qlist):
    code = ql.code_answers(_record(), qlist, printed_pages={2: "1", 3: "2"})
    assert set(code) == {q.key for q in qlist if q.answered_by != "model"}
    # Pages are the record's 0-based index plus one.
    assert code["plan_sheet_pages"].value == [5]
    # A bound report's kind takes the row's option nearest the record's
    # word, and none when none is near.
    studies = code["attached_studies"]
    assert studies.value == [
        {"title": "An earlier study", "kind": "appended study",
         "pdf_pages": [9]},
        {"title": "Some logs", "kind": None, "pdf_pages": [10]}]
    assert "the record's 'data_report' is not near one option" in \
        studies.note
    pages = code["page_kinds"].value
    assert pages[0] == {"pdf_page": 1, "sheet_class": "cover",
                        "printed_page": ""}
    assert pages[2]["printed_page"] == "1"
    # B-1 and its well B-01 are one hole; the hand auger is not counted.
    holes = code["hole_total"]
    assert (holes.value, holes.pdf_pages) == (2, [7, 8])
    assert "1 hand-auger log(s) not counted" in holes.note
    assert "left out" not in holes.note       # no log in the data report
    assert code["uses_piles"].value is True
    assert code["uses_piles"].pdf_pages == [9]
    assert "1.5 m (while drilling)" in code["water_seen"].value
    assert all(a.source == "code" for a in code.values())
    # A Both row the code map gives no reading.
    assert code["back_figure_total"].note == "not on the record"


def test_code_never_makes_up_a_value(qlist):
    code = ql.code_answers(_record(calculations=[], investigations=[],
                                   page_labels=[]), qlist)
    gap = code["back_figure_total"]
    assert (gap.status, gap.value, gap.note) == (
        "not said", None, "not on the record")
    # No pile calculation is not a no, and no log is not a zero.
    assert code["uses_piles"].status == "not said"
    assert code["uses_piles"].value is None
    assert code["hole_total"].status == "not said"
    assert code["plan_sheet_pages"].status == "not said"
    assert code["is_soil_study"].note == "not on the record"


def _counts(qlist, readings):
    """The list's hole_total row with a made-up prior-count row beside it,
    under ``readings``."""
    earlier = ql.Question(row=99, output="Earlier holes", category="Site",
                          data_type="integer", allow_null=True,
                          prompt="How many earlier holes?",
                          key="earlier_hole_total", answered_by="both")
    return ql.QuestionList(questions=[qlist.by_key()["hole_total"], earlier],
                           readings=readings, map_name="made-up")


def test_a_log_on_a_bound_prior_report_s_pages_is_a_prior_hole(qlist):
    record = _record(
        investigations=[_log("B-1", page=6), _log("B-2", page=7),
                        _log("P-9", page=8)],    # inside the prior's pages
        bound_documents=[BoundReport(
            bound_id="bound1", title="An earlier study",
            kind="appended_prior_report", first_page=8, last_page=9,
            n_pages=2)])
    child = ReportRecord(document=DocumentFacts(report_id="R00b", n_pages=2))
    rows = _counts(qlist, {"hole_total": ("hole_count", "boring,well"),
                           "earlier_hole_total": ("prior_hole_count",
                                                  "boring,well")})
    code = ql.code_answers(record, rows, bound={"bound1": child})
    assert (code["hole_total"].value, code["hole_total"].pdf_pages) == (
        2, [7, 8])
    assert (code["earlier_hole_total"].value,
            code["earlier_hole_total"].pdf_pages) == (1, [9])
    # The prior report's own record is merged by hole key, not added twice.
    child = ReportRecord(investigations=[_log("P-9", page=8),
                                         _log("P-10", page=9)])
    code = ql.code_answers(record, rows, bound={"bound1": child})
    assert code["earlier_hole_total"].value == 2
    assert code["hole_total"].value == 2


def test_a_log_on_a_bound_data_report_s_pages_counts_for_neither(qlist):
    record = _record(
        investigations=[_log("B-1", page=6),
                        _log("TP-5", kind="test_pit", page=9)],
        bound_documents=[BoundReport(bound_id="bound2", title="Some logs",
                                     kind="data_report", first_page=9,
                                     last_page=9, n_pages=1)])
    rows = _counts(qlist, {"hole_total": ("hole_count", "test_pit"),
                           "earlier_hole_total": ("prior_hole_count",
                                                  "test_pit")})
    code = ql.code_answers(record, rows)
    pits = code["hole_total"]
    assert (pits.status, pits.value) == ("not said", None)
    assert "1 test pit log(s) on the pages of a bound data report left " \
        "out" in pits.note
    assert code["earlier_hole_total"].status == "not said"
    assert "no prior report is bound in" in code["earlier_hole_total"].note


def test_a_checkbox_mark_is_not_a_chemical_result(qlist):
    row = ql.Question(row=99, output="Chemistry", category="Site",
                      data_type="string", allow_null=True,
                      prompt="Give the chemistry.", key="chemistry",
                      answered_by="both")
    rows = ql.QuestionList(questions=[row], map_name="made-up",
                           readings={"chemistry": ("chemical_results", "")})
    record = _record(lab_tests=[LabTest(
        kind="chemical", investigation_id="B-1", pages=[9],
        result={"kind": "chemical", "pH": 7.5,
                "sulfides": ":selected: Positive",
                "chloride": ":unselected:"})])
    said = ql.code_answers(record, rows)["chemistry"].value
    assert "pH 7.5" in said and "sulfides Positive" in said
    assert "selected" not in said and "chloride" not in said


def test_a_bound_volume_s_logs_count_when_its_record_is_given(qlist):
    record = _record(bound_documents=[BoundReport(
        bound_id="bound1", kind="volume", first_page=8, last_page=9,
        counts={"investigations": 1})])
    without = ql.code_answers(record, qlist)["hole_total"]
    assert without.status == "not said"
    assert "its record was not given" in without.note
    child = ReportRecord(investigations=[_log("B-9", page=8)])
    given = ql.code_answers(record, qlist, bound={"bound1": child})
    assert given["hole_total"].value == 3


# ---------------------------------------------------------------------------
# a Both row
# ---------------------------------------------------------------------------

def _q_of(qlist, key):
    return qlist.by_key()[key]


def test_compare_sets_code_beside_the_model(qlist):
    holes = _q_of(qlist, "hole_total")
    same = ql.compare(holes, Answer(value=2, status="answered"),
                      Answer(value=2, status="answered", source="code"))
    assert (same.source, same.value, same.code_value) == ("both: agree", 2, 2)
    other = ql.compare(holes, Answer(value=3, status="answered"),
                       Answer(value=2, status="answered"))
    assert (other.source, other.value, other.code_value) == (
        "both: disagree", 3, 2)
    assert "code: 2" in other.note
    features = _q_of(qlist, "nearby_features")
    assert ql.compare(features, Answer(value=["Creek ", "cliff"],
                                       status="answered"),
                      Answer(value=["cliff", "creek"], status="answered")
                      ).source == "both: agree"
    water = _q_of(qlist, "water_seen")
    assert ql.compare(water, Answer(value="Dry.", status="answered"),
                      Answer(value="B-1: 1.5 m", status="answered")
                      ).source == "both: compare by eye"


def test_records_compare_on_their_first_part(qlist):
    studies = _q_of(qlist, "attached_studies")
    model = Answer(value=[{"title": "An earlier study ", "kind": None,
                           "pdf_pages": [9, 10]}], status="answered")
    code = Answer(value=[{"title": "an earlier study", "kind":
                          "appended study", "pdf_pages": [9]}],
                  status="answered")
    assert ql.compare(studies, model, code).source == "both: agree"
    other = Answer(value=[{"title": "A later study", "kind": None,
                           "pdf_pages": [9]}], status="answered")
    assert ql.compare(studies, other, code).source == "both: disagree"
    # A Dictionary's first part is its key.
    colours = _q_of(qlist, "colour_names")
    assert ql.compare(colours,
                      Answer(value=[{"name": "A", "colour": "Red"}],
                             status="answered"),
                      Answer(value=[{"name": "a", "colour": "Blue"}],
                             status="answered")).source == "both: agree"


def test_compare_with_no_reading_on_one_side(qlist):
    holes = _q_of(qlist, "hole_total")
    none = ql.compare(holes, Answer(value=3, status="answered"),
                      Answer(status="not said", note="no boring log"))
    assert (none.source, none.value, none.code_value) == (
        "report text", 3, None)
    assert "code: no boring log" in none.note
    unsaid = ql.compare(holes, Answer(status="not said"),
                        Answer(value=2, status="answered"))
    assert (unsaid.source, unsaid.status) == ("both: disagree", "not said")


# ---------------------------------------------------------------------------
# the calls: screening, the gate, the main call
# ---------------------------------------------------------------------------

def _screening(soil=True, log_book=False, advice=False):
    return {"is_soil_study": _answer(soil, pages=[1]),
            "is_log_book": _answer(log_book),
            "is_advice_only": _answer(advice)}


def _main(**changes):
    out = {k: v for k, v in _full().items()
           if k not in _screening() and k not in (
               "attached_studies", "plan_sheet_pages", "page_kinds")}
    out.update(changes)
    return out


def _main_for_a_log_book():
    return {k: v for k, v in _main().items()
            if k not in ("post_heights", "uses_piles")}


PAGES = {0: "A made-up cover.", 1: "A made-up letter.", 2: "Some prose."}


def test_a_full_run_on_a_fake_engine(qlist):
    engine = FakeEngine([{"final": _screening()}, {"final": _main()}])
    run = ql.run_questions(_record(), PAGES, qlist, engine)
    assert engine.n_calls == 2 and run.closed == {}
    screening, main = (c["messages"][0]["content"][0]["text"]
                       for c in engine.calls)
    assert "=== PDF page 1 ===\nA made-up cover." in screening
    assert "Quote numbers with their units." in screening   # the rules
    assert "PDF page 5: plan" in main
    # Code's readings never go with a call: a Both row's two readings are
    # made apart.
    for text in (screening, main):
        assert "hole_total" not in text and "1.5 m" not in text
        assert "CODE'S" not in text and "code's" not in text
    assert "- is_soil_study: yes" in main           # the screening answers
    # The two calls offer only the rows a model answers.
    asked = set(engine.calls[1]["output_format"].model_fields)
    assert "hole_total" in asked and "attached_studies" not in asked
    assert "page_kinds" not in asked
    answers = run.answers
    assert answers["attached_studies"].source == "code"
    assert answers["hole_total"].source == "both: disagree"   # 3 against 2
    assert answers["well_count"].value == 2
    assert answers["site_flat"].status == "not said"
    assert answers["uses_piles"].status == "does not apply"
    assert answers["back_figure_total"].source == "report text"
    assert [c["call"] for c in run.calls] == ["screening", "main"]


def test_a_gate_closing_every_row_answers_screening_only(qlist):
    engine = FakeEngine([{"final": _screening(soil=False)}])
    run = ql.run_questions(_record(), PAGES, qlist, engine)
    assert engine.n_calls == 1
    for q in ql.main_questions(qlist):
        answer = run.answers[q.key]
        assert (answer.status, answer.source) == ("does not apply",
                                                  "screening"), q.key
    # Code's own reading is kept in the note of a row the gate closed.
    assert "code's reading: yes" in run.answers["uses_piles"].note
    assert run.answers["attached_studies"].source == "code"
    assert run.answers["is_soil_study"].value is False


def test_a_gate_closing_one_category(qlist):
    engine = FakeEngine([{"final": _screening(log_book=True)},
                         {"final": _main_for_a_log_book()}])
    run = ql.run_questions(_record(), PAGES, qlist, engine)
    closed = {q.key for q in qlist if q.category == "Advice"}
    assert set(run.closed) == closed
    asked = set(engine.calls[1]["output_format"].model_fields)
    assert not asked & closed and "site_flat" in asked
    assert run.answers["post_heights"].status == "does not apply"
    assert run.answers["well_count"].status == "answered"


def test_contradictory_screening_is_flagged_and_closes_nothing(qlist):
    # Log book alone would close the Advice rows; with advice-only also yes
    # the screening contradicts itself, so every row is asked.
    engine = FakeEngine([{"final": _screening(log_book=True, advice=True)},
                         {"final": _main()}])
    run = ql.run_questions(_record(), PAGES, qlist, engine)
    for key in ("is_log_book", "is_advice_only"):
        assert "both answered yes" in run.answers[key].note
    assert run.closed == {}
    asked = set(engine.calls[1]["output_format"].model_fields)
    assert {"post_heights", "uses_piles"} <= asked
    assert run.answers["post_heights"].status == "answered"
    assert any("nothing is closed and every row is asked" in w
               for w in run.warnings)


def test_not_said_and_does_not_apply_stay_apart(qlist):
    engine = FakeEngine([{"final": _screening()}, {"final": _main(
        site_flat=_answer(None, "not said"),
        well_count=_answer(None, "does not apply"))}])
    run = ql.run_questions(_record(), PAGES, qlist, engine)
    assert run.answers["site_flat"].status == "not said"
    assert run.answers["well_count"].status == "does not apply"


def test_what_does_not_hold_up_in_an_answer_is_noted(qlist):
    engine = FakeEngine([{"final": _screening()}, {"final": _main(
        well_count=_answer(2, pages=[4, 99]),
        flood_risk=_answer("low", status="not said"))}])
    run = ql.run_questions(_record(), PAGES, qlist, engine)
    wells = run.answers["well_count"]
    assert wells.pdf_pages == [4] and "99" in wells.note
    risk = run.answers["flood_risk"]
    assert risk.value is None and "a value came with status" in risk.note
    # Allow Null is off on the issue month: a blank there is noted.
    engine = FakeEngine([{"final": _screening()}, {"final": _main(
        issue_month=_answer(None, "not said"))}])
    run = ql.run_questions(_record(), PAGES, qlist, engine)
    assert "does not expect this row blank" in run.answers["issue_month"].note


def test_a_call_with_no_answer_leaves_its_rows_unanswered(qlist):
    engine = FakeEngine([{"final": _screening()},
                         {"text": "cut", "stop_reason": "max_tokens"}])
    run = ql.run_questions(_record(), PAGES, qlist, engine)
    assert run.answers["site_flat"].status == ""
    assert "no answer came back" in run.answers["site_flat"].note
    holes = run.answers["hole_total"]
    assert (holes.status, holes.source) == ("", "")
    assert "code: 2" in holes.note
    assert any("main call gave no structured answer" in w
               for w in run.warnings)
    assert run.calls[1]["answered"] is False


def test_the_pages_read_follow_the_narrative_reader_s_rule():
    pages = ql.pages_to_read(_record())
    assert pages == list(range(10))          # a short file is all front
    lines = ql.page_list_text(_record(), {2: "1", 3: "2"}).splitlines()
    assert lines[:3] == ["PDF page 1: cover", "PDF page 2: letter",
                         "PDF pages 3-4 (printed 1-2): narrative"]


# ---------------------------------------------------------------------------
# the answer sheet
# ---------------------------------------------------------------------------

def test_the_answer_sheet_round_trips(qlist, tmp_path):
    import openpyxl

    engine = FakeEngine([{"final": _screening()}, {"final": _main(
        nearby_features=_answer(["creek", "cliff"], pages=[3, 4, 5, 9]))}])
    record = _record(page_labels=[PageLabel(page=n, label="lab_test")
                                  for n in range(ql.MAX_CELL_RECORDS + 5)])
    record.document.n_pages = ql.MAX_CELL_RECORDS + 5
    run = ql.run_questions(record, PAGES, qlist, engine)
    path = ql.answer_sheet(qlist, run.answers, tmp_path / "R00.xlsx",
                           report_id="R00", calls=run.calls)
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["Answers", "Page kinds", "Run"]
    rows = list(book["Answers"].iter_rows(values_only=True))
    assert list(rows[0]) == list(ql.SHEET_COLUMNS)
    by_key = {r[3]: r for r in rows[1:]}
    assert len(by_key) == len(ROWS)
    assert by_key["nearby_features"][4] == "creek; cliff"
    assert by_key["nearby_features"][5] == "3-5, 9"
    assert by_key["site_flat"][4] is None             # not said: empty
    assert by_key["uses_piles"][4] == "N/A"           # does not apply
    assert by_key["colour_names"][4] == "A: Red"
    assert by_key["post_heights"][4] == (
        "post: A; height text: 2 m; finish: bare; pdf page: 3")
    assert by_key["attached_studies"][4] == (
        "title: An earlier study; kind: appended study; pdf pages: 9\n"
        "title: Some logs; pdf pages: 10")
    assert by_key["hole_total"][6] == "both: disagree"
    assert by_key["site_flat"][8] is None and by_key["site_flat"][9] is None
    # A long list of records goes on its own tab, named by its Output.
    assert "on the 'Page kinds' tab" in by_key["page_kinds"][4]
    tab = list(book["Page kinds"].iter_rows(values_only=True))
    assert tab[0] == ("pdf_page", "sheet_class", "printed_page")
    assert tab[1][:2] == (1, "lab_test") and len(tab) == len(
        record.page_labels) + 1
    info = {r[0]: r[1] for r in book["Run"].iter_rows(values_only=True)
            if r and r[0]}
    assert info["Report"] == "R00"
    assert info["Question list"] == "questions.xlsx"
    assert info["Code map"] == "questions.code.json"
    assert info["List modified"]
    assert info["screening"] == "fake-model" and info["main"] == "fake-model"
