"""The owner's report question list, as the schema a report is read against
(BACKLOG RI-90).

THE LIST is a workbook the owner keeps: a ``Rules`` sheet (the conventions
sent with every call: a heading in its first row, then one rule a row) and
a sheet of questions, one a row, under the headings Output, Category, Data
Type, Allow Null, Other Validation, Prompt, Key and Answered by. It names
the owner's organisation and stays private: it is read from a path given at
run time (:func:`load_questions`, or the :data:`LIST_ENV` variable), and
nothing in this package holds a question, a key or a pick list. Every key,
prompt and closed list below comes from the file.

THE CODE MAP is a small JSON file kept beside the list (the list's name with
:data:`CODE_MAP_SUFFIX`), as private as the list because it names its keys.
It is the one place that says which of code's :data:`READINGS` answers each
"Code" or "Both" row, which screening answers close other rows
(:class:`Gate`), and which pairs of screening answers cannot both be yes::

    {"readings": {"<Key>": "hole_count:boring,well", ...},
     "gates": [{"key": "<Key>", "closes_when": false, "closes": "*"},
               {"key": "<Key>", "closes_when": true, "closes": "<Category>"}],
     "not_both_yes": [["<Key>", "<Key>"]]}

WHAT IS BUILT FROM IT.

* :func:`answer_schema`: a pydantic model with one field per Key, each an
  answer of four parts: ``value``, typed from the Data Type (a closed list
  "[a, b]" is a Literal, "parts: x; y [p, q]" a record); ``status``;
  ``pdf_pages``; ``note``. A Dictionary (parts marked "(key)" and
  "(value)") comes back as a list of records, one per key, because strict
  structured output carries no open dictionary.
* ``status`` keeps the owner's two blanks apart: ``not said`` (the report
  is silent; an empty cell on the sheet) and ``does not apply`` (N/A). One
  is never folded into the other.
* :func:`code_answers`: the rows answered by "Code" or "Both, compared",
  from a finished report record; :func:`compare` sets code's reading beside
  the model's on a "Both" row.
* :func:`run_questions`: a screening call, then one call for the rest
  (the code map's gates say which screening answers close which rows).
* :func:`answer_sheet`: the answers in the owner's layout, with empty
  "Right?" and "Correct answer" columns for the owner to mark.

WHAT CODE ANSWERS. Only what the record shows: logs found, a calculation of
a kind, a caption count, the pages labelled a plan. The absence of a log or
a calculation is not taken as a no or a zero, because a report can describe
what it does not bind in: such a row stays ``not said``, with the reason in
its note. A row the code map gives no reading says "not on the record".

PAGES. Every page an answer gives is a PDF page numbered from 1 at the
file's first page, as a PDF viewer shows it. The record keeps 0-based
indexes, so code adds one.

Nothing on today's path imports this module; the graph wiring is a later
item (RI-90, after RI-87 piece 8).
"""

from __future__ import annotations

import os
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from typing import (Any, Callable, Dict, Iterator, List, Literal, Mapping,
                    Optional, Sequence, Tuple, Union, get_args)

from pydantic import (BaseModel, ConfigDict, Field, ValidationError,
                      create_model)

__all__ = [
    "LIST_ENV", "CODE_MAP_SUFFIX", "DATA_TYPES", "STATUSES", "READINGS",
    "QuestionListError", "Part", "Question", "Gate", "QuestionList", "Answer",
    "QuestionRun", "load_questions", "answer_schema", "screening_questions",
    "main_questions", "closed_rows", "code_answers", "compare",
    "pages_to_read", "page_texts", "page_list_text", "run_questions",
    "answer_sheet",
]

#: The variable that names the list's path when no path is given.
LIST_ENV = "GSE_REPORT_QUESTIONS"

#: The questions sheet's headings. The first six must be there; Allow Null
#: and Other Validation may be left out. Any other column is ignored.
REQUIRED_COLUMNS: Tuple[str, ...] = (
    "Output", "Category", "Data Type", "Prompt", "Key", "Answered by")
OPTIONAL_COLUMNS: Tuple[str, ...] = ("Allow Null", "Other Validation")

#: The Data Types the list may use, matched without regard to case.
DATA_TYPES: Tuple[str, ...] = (
    "string", "boolean", "integer", "date", "list of strings",
    "list of records", "dictionary", "integer or list of integers")

#: "Answered by", by its first word: "Code", "Both, compared", "Report text
#: (model)".
ANSWERED_BY: Dict[str, str] = {"code": "code", "both": "both",
                               "report": "model"}

#: What a model's answer says about the report.
STATUSES: Tuple[str, ...] = ("answered", "not said", "does not apply")
Status = Literal["answered", "not said", "does not apply"]

#: The category whose rows are asked first.
SCREENING = "screening"

#: The code map's file, beside the list: ``<list name><suffix>``.
CODE_MAP_SUFFIX = ".code.json"

#: What the "From" column says.
FROM_CODE = "code"
FROM_TEXT = "report text"
FROM_AGREE = "both: agree"
FROM_DISAGREE = "both: disagree"
#: A written description beside code's list: no rule compares the two.
FROM_BY_EYE = "both: compare by eye"
FROM_SCREENING = "screening"

#: The version of what the calls ask, kept on the Run tab.
PROMPT_VERSION = "ql-1"

#: How many of the pages read go to the screening call: the first ones,
#: where the cover, the letter and the contents are.
SCREENING_PAGES = 8

#: A list of records longer than this goes on its own tab of the sheet.
MAX_CELL_RECORDS = 30
#: How many holes or samples a written cross-check lists before it stops.
MAX_LISTED = 40


class QuestionListError(ValueError):
    """The list cannot be read: every malformed row, one line each."""


# ---------------------------------------------------------------------------
# the list
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Part:
    """One part of a record: its name, its closed list if it has one, and
    "key" or "value" on a Dictionary's parts."""

    name: str
    options: Tuple[str, ...] = ()
    role: str = ""


@dataclass(frozen=True)
class Question:
    """One row of the list."""

    row: int
    output: str
    category: str
    data_type: str
    allow_null: bool
    prompt: str
    key: str
    answered_by: str
    options: Tuple[str, ...] = ()
    parts: Tuple[Part, ...] = ()
    #: Other Validation as written when it is neither a closed list nor
    #: parts: a Date's format, or a pick list held elsewhere.
    validation: str = ""

    @property
    def screening(self) -> bool:
        return self.category.lower() == SCREENING


@dataclass(frozen=True)
class Gate:
    """A screening answer that closes other rows: when the Boolean row
    ``key`` is answered ``closes_when``, the rows of ``category`` do not
    apply ("*" closes every row outside Screening)."""

    key: str
    closes_when: bool
    category: str = "*"


@dataclass
class QuestionList:
    """The list as read: its rows, its rules and the file they came from,
    with its code map: code's reading for each Code or Both row (key ->
    reading name and argument), the gates, and the pairs of screening
    answers that cannot both be yes."""

    questions: List[Question]
    rules: List[str] = field(default_factory=list)
    name: str = ""
    modified: Optional[datetime] = None
    readings: Dict[str, Tuple[str, str]] = field(default_factory=dict)
    gates: Tuple[Gate, ...] = ()
    not_both_yes: Tuple[Tuple[str, str], ...] = ()
    map_name: str = ""

    def __iter__(self) -> Iterator[Question]:
        return iter(self.questions)

    def __len__(self) -> int:
        return len(self.questions)

    def by_key(self) -> Dict[str, Question]:
        return {q.key: q for q in self.questions}


_RE_KEY = re.compile(r"^[a-z][a-z0-9_]*$")
_RE_LIST = re.compile(r"^\[(.*)\]$", re.S)
_RE_PART = re.compile(
    r"^([A-Za-z_][A-Za-z0-9_]*)\s*(?:\((key|value)\))?\s*(?:\[(.*)\])?$",
    re.S)


def _cell(value: Any) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _options(text: str) -> Tuple[str, ...]:
    return tuple(o for o in (" ".join(p.split()) for p in text.split(","))
                 if o)


def _flag(value: Any) -> Optional[bool]:
    """Allow Null as written: blank is no; None when it cannot be read."""
    if value is None or isinstance(value, bool):
        return bool(value)
    text = _cell(value).lower()
    if text in ("", "no", "n", "false", "0"):
        return False
    if text in ("yes", "y", "true", "1"):
        return True
    return None


def _parse_parts(text: str) -> Tuple[Tuple[Part, ...], str]:
    """``parts: x; y [p, q]; z (key)`` as parts, and what is wrong ("" when
    nothing is)."""
    parts: List[Part] = []
    for piece in text.split(":", 1)[1].split(";"):
        piece = " ".join(piece.split())
        if not piece:
            continue
        match = _RE_PART.match(piece)
        if match is None:
            return (), f"the part {piece!r} is not a name with an optional " \
                       f"'(key)', '(value)' or closed list"
        options = _options(match.group(3)) if match.group(3) else ()
        if match.group(3) is not None and not options:
            return (), f"the part {match.group(1)!r} has an empty list"
        parts.append(Part(match.group(1).lower(), options,
                          match.group(2) or ""))
    names = [p.name for p in parts]
    if not parts:
        return (), "'parts:' names no part"
    if len(set(names)) != len(names):
        return (), "a part is named twice"
    roles = [p.role for p in parts if p.role]
    if roles and sorted(roles) != ["key", "value"]:
        return (), "a Dictionary marks one part '(key)' and one '(value)'"
    # The key part first, so a record reads "key: value".
    parts.sort(key=lambda p: {"key": 0, "value": 1}.get(p.role, 2))
    return tuple(parts), ""


def _question(row: int, cells: Mapping[str, Any]
              ) -> Tuple[Optional[Question], List[str]]:
    """One row as a question, or the reasons it is malformed."""
    key = _cell(cells.get("Key"))
    where = f"row {row}" + (f" (Key {key!r})" if key else "")
    problems: List[str] = []
    if not _RE_KEY.match(key):
        problems.append("its Key must be lower-case letters, digits and "
                        "underscores, starting with a letter")
    data_type = _cell(cells.get("Data Type")).lower()
    if data_type not in DATA_TYPES:
        problems.append(f"its Data Type {_cell(cells.get('Data Type'))!r} "
                        f"is not one of: " + ", ".join(DATA_TYPES))
    answered = ANSWERED_BY.get(
        (_cell(cells.get("Answered by")).lower().split() or [""])[0]
        .rstrip(","))
    if answered is None:
        problems.append(f"its Answered by {_cell(cells.get('Answered by'))!r}"
                        f" is not Code, Both, compared or Report text "
                        f"(model)")
    prompt = _cell(cells.get("Prompt"))
    if not prompt:
        problems.append("it has no Prompt")
    allow_null = _flag(cells.get("Allow Null"))
    if allow_null is None:
        problems.append(f"its Allow Null {_cell(cells.get('Allow Null'))!r} "
                        f"is not yes, no or blank")
    text = _cell(cells.get("Other Validation"))
    options: Tuple[str, ...] = ()
    parts: Tuple[Part, ...] = ()
    validation = ""
    listed = _RE_LIST.match(text)
    if text.lower().startswith("parts:"):
        parts, wrong = _parse_parts(text)
        if wrong:
            problems.append(f"its Other Validation: {wrong}")
        elif data_type not in ("list of records", "dictionary"):
            problems.append("its Other Validation gives parts, which only a "
                            "List of records or a Dictionary has")
    elif listed and data_type in ("string", "list of strings"):
        options = _options(listed.group(1))
        if not options:
            problems.append("its Other Validation is an empty list")
    elif listed and data_type == "date":
        validation = listed.group(1).strip()        # the date's format
    elif listed:
        problems.append("its Other Validation is a closed list, which only "
                        "a String or a List of strings has")
    else:
        validation = text
    if data_type in ("list of records", "dictionary") and not parts \
            and not problems:
        problems.append(f"a {data_type.title()} needs its parts under Other "
                        f"Validation ('parts: a; b')")
    if problems:
        return None, [f"{where}: {p}" for p in problems]
    return Question(
        row=row, output=_cell(cells.get("Output")),
        category=_cell(cells.get("Category")), data_type=data_type,
        allow_null=bool(allow_null), prompt=prompt, key=key,
        answered_by=answered or "", options=options, parts=parts,
        validation=validation), []


def _list_path(path: Union[str, Path, None]) -> Path:
    if path is None:
        named = os.environ.get(LIST_ENV, "").strip()
        if not named:
            raise QuestionListError(
                f"no question list: give its path, or set {LIST_ENV}")
        path = named
    out = Path(path).expanduser()
    if not out.is_file():
        raise QuestionListError(f"the question list is not a file: {out}")
    return out


def _map_path(list_path: Path, code_map: Union[str, Path, None]
              ) -> Optional[Path]:
    if code_map is not None:
        out = Path(code_map).expanduser()
        if not out.is_file():
            raise QuestionListError(f"the code map is not a file: {out}")
        return out
    beside = list_path.with_name(list_path.stem + CODE_MAP_SUFFIX)
    return beside if beside.is_file() else None


def _read_map(path: Path, questions: Sequence[Question]
              ) -> Tuple[Dict[str, Tuple[str, str]], Tuple[Gate, ...],
                         Tuple[Tuple[str, str], ...], List[str]]:
    """The code map at ``path``, checked against the list: its readings,
    gates and pairs, and what is wrong with it, one line each."""
    import json

    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        return {}, (), (), [f"the code map {path.name} is not JSON: {exc}"]
    if not isinstance(data, dict):
        return {}, (), (), [f"the code map {path.name} is not an object"]
    by_key = {q.key: q for q in questions}
    categories = {q.category.lower() for q in questions}
    problems: List[str] = []
    readings: Dict[str, Tuple[str, str]] = {}
    for key, said in dict(data.get("readings") or {}).items():
        name, _, arg = str(said).partition(":")
        name, arg = name.strip(), arg.strip()
        where = f"the code map's reading for Key {key!r}"
        if key not in by_key:
            problems.append(f"{where}: the list has no such Key")
        elif by_key[key].answered_by == "model":
            problems.append(f"{where}: the list gives that row to the model")
        elif name not in READINGS:
            problems.append(f"{where}: {name!r} is not one of: "
                            + ", ".join(READINGS))
        else:
            try:
                READINGS[name](arg)
            except ValueError as exc:
                problems.append(f"{where}: {exc}")
            else:
                readings[key] = (name, arg)
    gates: List[Gate] = []
    for gate in data.get("gates") or []:
        gate = gate if isinstance(gate, dict) else {}
        key = str(gate.get("key", ""))
        closes = str(gate.get("closes", "*")).strip() or "*"
        q = by_key.get(key)
        where = f"the code map's gate on Key {key!r}"
        if q is None or not (q.screening and q.data_type == "boolean"):
            problems.append(f"{where}: a gate is a Boolean row in the "
                            f"Screening category")
        elif q.answered_by == "code":
            problems.append(f"{where}: the list gives that row to code, and "
                            f"only an answer screening asks can close rows")
        elif not isinstance(gate.get("closes_when"), bool):
            problems.append(f"{where}: closes_when is true or false")
        elif closes != "*" and closes.lower() not in categories:
            problems.append(f"{where}: no row is in the category {closes!r}")
        else:
            gates.append(Gate(key, gate["closes_when"], closes))
    pairs: List[Tuple[str, str]] = []
    for pair in data.get("not_both_yes") or []:
        keys = tuple(str(k) for k in pair) if isinstance(
            pair, (list, tuple)) else ()
        if len(keys) != 2 or not all(k in by_key and by_key[k].data_type
                                     == "boolean" for k in keys):
            problems.append(f"the code map's pair {list(keys)!r}: a pair is "
                            f"two Boolean Keys of the list")
        elif any(by_key[k].answered_by == "code" for k in keys):
            problems.append(f"the code map's pair {list(keys)!r}: the list "
                            f"gives a row of it to code, and a pair is two "
                            f"answers screening asks")
        else:
            pairs.append((keys[0], keys[1]))
    return readings, tuple(gates), tuple(pairs), problems


def load_questions(path: Union[str, Path, None] = None, *,
                   code_map: Union[str, Path, None] = None) -> QuestionList:
    """The list at ``path`` (default: the :data:`LIST_ENV` variable): its
    rows and its rules, with its code map (``code_map``, default: the file
    beside the list named with :data:`CODE_MAP_SUFFIX`, if there is one).

    The questions sheet is the first whose first row has the headings
    :data:`REQUIRED_COLUMNS`; the rules come from a sheet named "Rules", if
    there is one. A malformed row, a Key used twice and a code map entry
    that does not fit the list are each named, every one at once, in the
    :class:`QuestionListError`. With no code map, screening closes no row
    and code reads no row.
    """
    import openpyxl

    path = _list_path(path)
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rules: List[str] = []
        sheet = None
        heads: Dict[str, int] = {}
        for ws in book.worksheets:
            rows = ws.iter_rows(values_only=True)
            first = next(rows, ()) or ()
            if ws.title.strip().lower() == "rules":
                rules = [_cell(r[0]) for r in rows if r and _cell(r[0])]
                continue
            found = {_cell(v).lower(): n for n, v in enumerate(first)
                     if _cell(v)}
            if sheet is None and all(c.lower() in found
                                     for c in REQUIRED_COLUMNS):
                sheet, heads = ws, found
        if sheet is None:
            raise QuestionListError(
                "no sheet has the headings " + ", ".join(REQUIRED_COLUMNS))
        questions: List[Question] = []
        problems: List[str] = []
        names = REQUIRED_COLUMNS + OPTIONAL_COLUMNS
        for row, values in enumerate(
                sheet.iter_rows(min_row=2, values_only=True), start=2):
            cells = {c: (values[heads[c.lower()]]
                         if c.lower() in heads
                         and heads[c.lower()] < len(values) else None)
                     for c in names}
            if not any(_cell(v) for v in cells.values()):
                continue
            question, wrong = _question(row, cells)
            problems.extend(wrong)
            if question is not None:
                questions.append(question)
    finally:
        book.close()
    seen: Dict[str, int] = {}
    for q in questions:
        if q.key in seen:
            problems.append(f"row {q.row} (Key {q.key!r}): the Key is "
                            f"already used on row {seen[q.key]}")
        seen.setdefault(q.key, q.row)
    mapped = _map_path(path, code_map)
    readings: Dict[str, Tuple[str, str]] = {}
    gates: Tuple[Gate, ...] = ()
    pairs: Tuple[Tuple[str, str], ...] = ()
    if mapped is not None:
        readings, gates, pairs, wrong = _read_map(mapped, questions)
        problems.extend(wrong)
    if problems:
        raise QuestionListError(
            f"the question list {path.name} has {len(problems)} problem(s):\n"
            + "\n".join(problems))
    return QuestionList(
        questions=questions, rules=rules, name=path.name,
        modified=datetime.fromtimestamp(path.stat().st_mtime),
        readings=readings, gates=gates, not_both_yes=pairs,
        map_name=mapped.name if mapped is not None else "")


def screening_questions(questions: Sequence[Question]) -> List[Question]:
    """The rows asked first: the Screening category."""
    return [q for q in questions if q.screening]


def main_questions(questions: Sequence[Question]) -> List[Question]:
    """Every other row."""
    return [q for q in questions if not q.screening]


# ---------------------------------------------------------------------------
# the answer schema
# ---------------------------------------------------------------------------

def _camel(key: str) -> str:
    return "".join(w[:1].upper() + w[1:] for w in key.split("_") if w)


def _part_type(part: Part) -> Any:
    """A part's type. Text as printed, but for a closed list and the PDF
    page parts (named ``pdf_page`` or ``pdf_pages``, or ending so), which
    are numbers. A part not printed is "" or null."""
    if part.options:
        return Optional[Literal[part.options]]           # type: ignore
    if part.name == "pdf_pages" or part.name.endswith("_pdf_pages"):
        return List[int]
    if part.name == "pdf_page" or part.name.endswith("_pdf_page"):
        return Optional[int]
    return str


_PART_WORDS = {
    "key": "the key", "value": "its value",
}


@lru_cache(maxsize=None)
def _record_model(question: Question) -> type:
    fields: Dict[str, Any] = {}
    for part in question.parts:
        said = part.name.replace("_", " ")
        if part.options:
            said += "; one of the listed values, or null if not printed"
        elif _part_type(part) is str:
            said += "; as printed, or empty if not printed"
        else:
            said += "; PDF page numbers, counting the file's first page as 1"
        if part.role:
            said = f"{_PART_WORDS[part.role]}: {said}"
        fields[part.name] = (_part_type(part), Field(description=said))
    return create_model(f"{_camel(question.key)}Item",
                        __config__=ConfigDict(extra="forbid"), **fields)


def _value_type(question: Question) -> Any:
    kind = question.data_type
    if kind == "boolean":
        return bool
    if kind == "integer":
        return int
    if kind == "integer or list of integers":
        return List[int]
    if kind in ("list of records", "dictionary"):
        return List[_record_model(question)]              # type: ignore
    if kind == "list of strings":
        return (List[Literal[question.options]]            # type: ignore
                if question.options else List[str])
    if question.options:                                   # string
        return Literal[question.options]                   # type: ignore
    return str                                             # string, date


def _value_words(question: Question) -> str:
    kind = question.data_type
    if kind == "boolean":
        said = "true for yes, false for no"
    elif kind == "integer":
        said = "a whole number"
    elif kind == "integer or list of integers":
        said = "a list of whole numbers; one number is a list of one"
    elif kind == "dictionary":
        said = "a list of records, one per key"
    elif kind == "list of records":
        said = "a list of records; an empty list if there are none"
    elif question.options:
        said = ("one or more of the listed values" if kind == "list of "
                "strings" else "one of the listed values")
    elif kind == "date":
        said = (f"the date as {question.validation}" if question.validation
                else "the date as printed")
    elif kind == "list of strings":
        said = "a list of text, each as the report words it"
    else:
        said = "text"
    return said


@lru_cache(maxsize=None)
def _answer_model(question: Question) -> type:
    # The question goes on its own answer model, as that model's
    # description, so each top-level property is a plain $ref with nothing
    # beside it. status, pdf_pages and note are the same on every row; the
    # call's instruction says what they hold once (_HOW_TO_ANSWER) rather
    # than 72 times over in the schema.
    return create_model(
        f"{_camel(question.key)}Answer",
        __config__=ConfigDict(extra="forbid"),
        __doc__=_question_words(question),
        value=(Optional[_value_type(question)],
               Field(default=None, description=_value_words(question))),
        status=(Status, ...), pdf_pages=(List[int], ...), note=(str, ...))


def _question_words(question: Question) -> str:
    said = f"{question.output}: {question.prompt}"
    if question.validation and question.data_type != "date":
        said += f" (Other validation: {question.validation})"
    return said


def answer_schema(questions: Sequence[Question],
                  name: str = "QuestionAnswers") -> type:
    """The answer model for ``questions``: one field per Key, in the list's
    order, each an answer (``value``, ``status``, ``pdf_pages``, ``note``)
    whose description is its Output and Prompt. Accepted by strict
    structured output: every object closed, no open dictionary, and each
    property a plain $ref."""
    fields = {q.key: (_answer_model(q), ...) for q in questions}
    return create_model(
        name, __config__=ConfigDict(extra="forbid"),
        __doc__="The answers to the report questions, one per key.",
        **fields)


# ---------------------------------------------------------------------------
# an answer, from code or from a model
# ---------------------------------------------------------------------------

@dataclass
class Answer:
    """One row's answer as the sheet shows it.

    ``status`` is one of :data:`STATUSES`, or "" when no answer came back at
    all (a call that failed). ``source`` is the sheet's From column.
    ``code_value`` is code's reading on a "Both" row, kept beside the
    model's."""

    value: Any = None
    status: str = "not said"
    pdf_pages: List[int] = field(default_factory=list)
    note: str = ""
    source: str = ""
    code_value: Any = None

    def noted(self, said: str) -> "Answer":
        if said:
            self.note = "; ".join(n for n in (self.note.strip(), said) if n)
        return self


def _plain(value: Any) -> Any:
    """A parsed value as plain data: a record as a dict, in part order."""
    if isinstance(value, BaseModel):
        return value.model_dump()
    if isinstance(value, list):
        return [_plain(v) for v in value]
    return value


def _fit(question: Question, answer: Answer) -> Answer:
    """``answer`` checked against the row's answer model: a reading that
    does not fit the row's Data Type is not given (``not said``, with why)."""
    try:
        parsed = _answer_model(question).model_validate({
            "value": answer.value, "status": answer.status or "not said",
            "pdf_pages": answer.pdf_pages, "note": answer.note})
    except ValidationError as exc:
        first = exc.errors()[0]
        return Answer(status="not said", source=answer.source, note=(
            f"code's reading does not fit the row's Data Type "
            f"({question.data_type}): {first.get('msg', '')}"))
    answer.value = _plain(parsed.value)
    return answer


# ---------------------------------------------------------------------------
# code's answers, from the record
# ---------------------------------------------------------------------------

@dataclass
class _Source:
    """What code reads: the record, its bound reports' own records (by
    bound_id) and the printed page numbers (0-based page -> as printed)."""

    record: Any
    bound: Dict[str, Any]
    printed: Dict[int, str]


def _as_record(record: Any) -> Any:
    from report_ingest.model import ReportRecord

    if isinstance(record, ReportRecord):
        return record
    return ReportRecord.model_validate(record)


def _given(value: Any, pages: Sequence[int] = (), note: str = "") -> Answer:
    return Answer(value=value, status="answered", source=FROM_CODE,
                  pdf_pages=sorted({int(p) for p in pages}), note=note)


def _not_given(why: str) -> Answer:
    return Answer(status="not said", source=FROM_CODE, note=why)


def _span(pages: Sequence[int]) -> str:
    """``[3, 4, 5, 9]`` as "3-5, 9"."""
    pages = sorted({int(p) for p in pages})
    runs: List[str] = []
    start = prev = None
    for page in pages + [None]:                           # type: ignore
        if page is not None and prev is not None and page == prev + 1:
            prev = page
            continue
        if start is not None:
            runs.append(str(start) if start == prev else f"{start}-{prev}")
        start = prev = page
    return ", ".join(runs)


def _words(text: str) -> set:
    return set(re.findall(r"[a-z0-9]+", str(text).lower()))


def _closest(options: Sequence[str], said: str) -> Optional[str]:
    """The listed option that shares the most words with ``said`` (the
    record's own word for a kind, "appended_prior_report"); None when none
    shares a word, or two share as many."""
    scored = sorted(((len(_words(o) & _words(said)), o) for o in options),
                    key=lambda s: -s[0])
    if not scored or not scored[0][0] or (
            len(scored) > 1 and scored[1][0] == scored[0][0]):
        return None
    return scored[0][1]


def _fill(question: Question, known: Mapping[str, Any],
          notes: List[str]) -> Dict[str, Any]:
    """One record in the row's parts, from what code read (``known``, by
    the record's own field names). A part takes the field of its name, else
    of the last or first word of its name ("sheet_class" takes "class"); a
    closed-list part takes the option nearest code's word (:func:`_closest`),
    and says so in ``notes`` when none is near; a part code has nothing for
    is left empty."""
    out: Dict[str, Any] = {}
    for part in question.parts:
        words = part.name.split("_")
        name = next((n for n in (part.name, words[-1], words[0])
                     if n in known), None)
        value = known.get(name) if name else None
        kind = _part_type(part)
        if part.options:
            word = value
            value = _closest(part.options, word) if word else None
            if word and value is None:
                notes.append(f"{part.name}: the record's {word!r} is not "
                             f"near one option of the row's list")
        elif kind is str:
            value = "" if value is None else str(value)
        elif value is None and part.name.endswith("pages"):
            value = []
        out[part.name] = value
    return out


def _bound_documents(src: _Source, q: Question) -> Answer:
    docs = list(src.record.bound_documents)
    notes: List[str] = []
    out: List[Dict[str, Any]] = []
    for doc in docs:
        pages = list(range(doc.first_page + 1, doc.last_page + 2))
        out.append(_fill(q, {"title": doc.title, "firm": doc.firm,
                             "date": doc.date, "kind": doc.kind,
                             "pdf_pages": pages, "pdf_page": pages[0]},
                         notes))
    return _given(out, [d.first_page + 1 for d in docs], "; ".join(
        ["the bound reports triage and the page labels found" if docs
         else "triage and the page labels found none"] + notes))


def _pages_labelled(label: str) -> Callable[[_Source, Question], Answer]:
    from report_ingest.corpus import LABELS

    if label not in LABELS:
        raise ValueError(f"the page label {label!r} is not one of: "
                         + ", ".join(LABELS))
    words = label.replace("_", " ")

    def read(src: _Source, q: Question) -> Answer:
        pages = [lab.page + 1 for lab in src.record.page_labels
                 if lab.label == label]
        if not pages:
            return _not_given(f"no page is labelled {words}")
        return _given(pages, pages, f"the pages labelled {words}; one may "
                                    f"be another {words} than the row asks "
                                    f"for")
    return read


def _page_labels(src: _Source, q: Question) -> Answer:
    labels = list(src.record.page_labels)
    if not labels:
        return _not_given("the record has no page labels")
    notes: List[str] = []
    out = [_fill(q, {"pdf_page": lab.page + 1, "pdf_pages": [lab.page + 1],
                     "label": lab.label, "class": lab.label,
                     "printed_page": src.printed.get(lab.page, "")}, notes)
           for lab in labels]
    return _given(out, (), "; ".join(
        ([] if src.printed else ["printed page numbers are not on the "
                                 "record"]) + notes[:3]))


def _bound_of(src: _Source, kinds: Sequence[str]
              ) -> Tuple[List[Any], List[str]]:
    """The records of the bound reports of these kinds, and a line for each
    one with logs whose record was not given."""
    records, missing = [], []
    for doc in src.record.bound_documents:
        if doc.kind not in kinds:
            continue
        child = src.bound.get(doc.bound_id)
        if child is not None:
            records.append(_as_record(child))
        elif int(doc.counts.get("investigations", 0) or 0):
            missing.append(f"the bound report on PDF pages "
                           f"{doc.first_page + 1}-{doc.last_page + 1} holds "
                           f"{doc.counts['investigations']} log(s), and its "
                           f"record was not given")
    return records, missing


def _holes(logs: Sequence[Any]) -> Tuple[int, List[int]]:
    """``(how many holes, the first PDF page of each)`` over these logs.
    One hole is one key (:func:`report_ingest.reconciler.hole_key`): a well
    logged in a boring is that boring; a log with no id counts alone; two
    holes printed under one name and kept apart by the reconciler (RI-135)
    are two."""
    from report_ingest.shared_names import hole_identity

    keys, pages = set(), []
    for i, inv in enumerate(logs):
        key = hole_identity(inv) or f"#{i}"
        if key not in keys and inv.pages:
            pages.append(inv.pages[0] + 1)
        keys.add(key)
    return len(keys), pages


#: Whose a main-record log is, by the kind of the bound report whose pages
#: hold its first page: a bound volume's (or no bound report's) is this
#: investigation's, a prior report's is the prior investigation's; a log in
#: any other bound report (a data report, another document) is neither.
_LOG_SCOPE = {"volume": "this", "appended_prior_report": "prior"}


def _placed(src: _Source
            ) -> Tuple[List[Any], List[Any], List[Tuple[str, Any]]]:
    """The main record's logs placed by the bound report whose
    first_page..last_page holds each log's first page: ``(this
    investigation's, the prior reports', the logs left out with the kind of
    the bound report holding each)``."""
    this, prior, left = [], [], []
    for inv in src.record.investigations:
        first = inv.pages[0] if inv.pages else None
        doc = next((d for d in src.record.bound_documents
                    if first is not None
                    and d.first_page <= first <= d.last_page), None)
        scope = "this" if doc is None else _LOG_SCOPE.get(doc.kind)
        if scope == "this":
            this.append(inv)
        elif scope == "prior":
            prior.append(inv)
        else:
            left.append((doc.kind or "other", inv))
    return this, prior, left


def _kinds(arg: str, allowed: Sequence[str]) -> Tuple[str, ...]:
    """``arg`` ("boring,well") as kinds, each one the record knows."""
    kinds = tuple(k.strip() for k in arg.split(",") if k.strip())
    wrong = [k for k in kinds if k not in allowed]
    if not kinds or wrong:
        raise ValueError("give the kinds, from: " + ", ".join(allowed)
                         + (f" (not {', '.join(wrong)})" if wrong else ""))
    return kinds


def _said(kinds: Sequence[str]) -> str:
    """Kinds in words: "boring or well"."""
    words = [k.replace("_", " ") for k in kinds]
    return (", ".join(words[:-1]) + " or " + words[-1] if len(words) > 1
            else words[0])


def _count(prior: bool) -> Callable[[str], Callable[[_Source, Question],
                                                    Answer]]:
    """How many holes of the given log kinds this investigation logged, or
    (``prior``) the prior investigations bound in. A log in the main record
    is placed by the bound report whose pages hold it (:func:`_placed`),
    then merged by hole key with the bound reports' own records."""
    def make(arg: str) -> Callable[[_Source, Question], Answer]:
        from report_ingest.model import InvestigationKind

        kinds = _kinds(arg, get_args(InvestigationKind))
        words = _said(kinds)

        def read(src: _Source, q: Question) -> Answer:
            this, before, left = _placed(src)
            if prior:
                if not any(d.kind == "appended_prior_report"
                           for d in src.record.bound_documents):
                    return _not_given("no prior report is bound in; prior "
                                      "investigations the text describes "
                                      "are not on any record")
                records, missing = _bound_of(src, ("appended_prior_report",))
                logs = before + [i for r in records for i in r.investigations]
                where = "the logs of the bound prior reports"
            else:
                records, missing = _bound_of(src, ("volume",))
                logs = this + [i for r in records for i in r.investigations]
                where = "the logs of this investigation" + (
                    " (with its bound volumes)" if records else "")
            if missing:
                return _not_given("; ".join(missing))
            n, pages = _holes([i for i in logs if i.kind in kinds])
            augers = sum(1 for i in logs if i.kind == "hand_auger")
            notes = []
            if "boring" in kinds and "hand_auger" not in kinds and augers:
                notes.append(f"{augers} hand-auger log(s) not counted; the "
                             f"report decides whether they are borings")
            if not prior:
                out = Counter(kind for kind, inv in left if inv.kind in kinds)
                notes.extend(
                    f"{k} {words} log(s) on the pages of a bound "
                    f"{kind.replace('_', ' ')} left out: whether it is this "
                    f"investigation's is not on the record"
                    for kind, k in sorted(out.items()))
            if not n:
                return _not_given("; ".join(
                    [f"no {words} log in {where}; a report can count holes "
                     f"it does not bind in"] + notes))
            return _given(n, pages, "; ".join(
                [f"{words} holes counted from {where}"] + notes))
        return read
    return make


def _bound_priors(src: _Source, q: Question) -> Answer:
    prior = [d for d in src.record.bound_documents
             if d.kind == "appended_prior_report"]
    if not prior:
        return _not_given("no prior report is bound in; prior investigations "
                          "the text describes are not on any record")
    return _given(len(prior), [d.first_page + 1 for d in prior],
                  "the prior reports bound in; prior investigations the "
                  "text only describes are not counted")


def _captions(kind: str) -> Callable[[_Source, Question], Answer]:
    if not kind:
        raise ValueError("give what the captions are of, as the record "
                         "counts them (tables, figures)")

    def read(src: _Source, q: Question) -> Answer:
        counted = src.record.narrative.counted
        if kind not in counted:
            return _not_given(f"no caption count of {kind} on the record")
        return _given(int(counted[kind]), (),
                      f"{kind} captions counted in the main body")
    return read


def _logs(src: _Source) -> List[Any]:
    records, _missing = _bound_of(src, ("volume",))
    return [inv for r in [src.record] + records for inv in r.investigations]


_RE_SOFT = re.compile(r"\bsoft\b", re.I)


def _soft(src: _Source, q: Question) -> Answer:
    holes: List[str] = []
    pages: List[int] = []
    for inv in _logs(src):
        layers = [lay for lay in inv.layers
                  if _RE_SOFT.search(f"{lay.consistency} {lay.description}")]
        if layers:
            holes.append(inv.investigation_id or "a log with no id")
            pages.extend(lay.prov.page + 1 if lay.prov else
                         (inv.pages[0] + 1 if inv.pages else 1)
                         for lay in layers[:1])
    if not holes:
        return _not_given("no log layer is described as soft; the text may "
                          "describe soft ground the logs do not")
    return _given(True, pages, f"layers described as soft on the logs of "
                               f"{len(holes)} hole(s): "
                               + ", ".join(holes[:MAX_LISTED]))


def _depth(quantity: Any) -> str:
    return f"{quantity.value:g} {quantity.unit}".strip()


_WHEN = {"while_drilling": "while drilling", "at_completion": "at completion",
         "not_encountered": "", "unknown": ""}


def _reading(level: Any) -> str:
    said = "not encountered" if level.depth is None else _depth(level.depth)
    when = (f"{level.hours:g} h after" if level.when == "after_hours"
            and level.hours is not None else _WHEN.get(level.when,
                                                       level.when))
    extra = ", ".join(w for w in (when, level.date) if w)
    return f"{said} ({extra})" if extra else said


def _water(src: _Source, q: Question) -> Answer:
    rows: List[str] = []
    pages: List[int] = []
    for inv in _logs(src):
        if not inv.water:
            continue
        rows.append(f"{inv.investigation_id or 'a log with no id'}: "
                    + ", ".join(_reading(w) for w in inv.water))
        pages.extend(w.prov.page + 1 for w in inv.water if w.prov)
    if not rows:
        return _not_given("no water reading on the logs")
    more = len(rows) - MAX_LISTED
    return _given("the logs' water readings: " + "; ".join(rows[:MAX_LISTED])
                  + (f"; and {more} more hole(s)" if more > 0 else ""), pages,
                  "code's reading off the logs, to check the description "
                  "against")


def _calcs(src: _Source, kinds: Sequence[str]) -> List[Any]:
    return [c for c in src.record.calculations if c.kind in kinds]


def _calc_pages(calcs: Sequence[Any]) -> List[int]:
    return [p + 1 for c in calcs for p in c.pages]


def _calc_kinds(arg: str) -> Tuple[str, ...]:
    from report_ingest.model import CALC_KINDS

    return _kinds(arg, CALC_KINDS)


def _has_calc(arg: str) -> Callable[[_Source, Question], Answer]:
    kinds = _calc_kinds(arg)
    words = _said(kinds)

    def read(src: _Source, q: Question) -> Answer:
        calcs = _calcs(src, kinds)
        if not calcs:
            return _not_given(f"no {words} calculation on the record; a "
                              f"report can give its results without the "
                              f"printouts")
        return _given(True, _calc_pages(calcs),
                      f"{len(calcs)} {words} calculation printout(s)")
    return read


_RE_SEISMIC = re.compile(
    r"seismic|pseudo[- ]?static|earthquake|newmark|\bkh\b", re.I)


def _seismic(arg: str) -> Callable[[_Source, Question], Answer]:
    kinds = _calc_kinds(arg)
    words = _said(kinds)

    def read(src: _Source, q: Question) -> Answer:
        calcs = [c for c in _calcs(src, kinds)
                 if _RE_SEISMIC.search(" ".join(
                     [c.method, c.subject, c.summary, c.program or ""]
                     + [f"{v.name} {v.text}"
                        for v in (*c.inputs, *c.results)]))]
        if not calcs:
            return _not_given(f"no {words} calculation on the record names "
                              f"a seismic case")
        return _given(True, _calc_pages(calcs),
                      f"{len(calcs)} {words} printout(s) name a seismic case "
                      f"(seismic, pseudo-static, kh or Newmark)")
    return read


def _calc_list(src: _Source, q: Question) -> Answer:
    calcs = list(src.record.calculations)
    if not calcs:
        return _not_given("no calculation printout on the record")
    by_kind: Dict[str, List[Any]] = {}
    for calc in calcs:
        by_kind.setdefault(calc.kind, []).append(calc)
    rows = []
    for kind, group in by_kind.items():
        methods = list(dict.fromkeys(c.method for c in group if c.method))
        rows.append(f"{kind.replace('_', ' ')}: {len(group)} printout(s)"
                    + (f" ({'; '.join(methods[:3])})" if methods else "")
                    + f", PDF pages {_span(_calc_pages(group))}")
    return _given("the calculation printouts: " + "; ".join(rows),
                  _calc_pages(calcs),
                  "code's list of the printouts, to check the summary "
                  "against")


def _printed_value(arg: str) -> Callable[[_Source, Question], Answer]:
    """The text a calculation printout prints under a name: ``arg`` is the
    name's wordings, "a b|c d"."""
    names = [n.split() for n in arg.split("|") if n.split()]
    if not names:
        raise ValueError("give the printed name's wordings, 'a b|c d'")
    pattern = re.compile("|".join(r"\s*".join(map(re.escape, n))
                                  for n in names), re.I)
    words = " or ".join(" ".join(n) for n in names)

    def read(src: _Source, q: Question) -> Answer:
        said: List[str] = []
        pages: List[int] = []
        for calc in src.record.calculations:
            for v in (*calc.inputs, *calc.results):
                if pattern.search(v.name) and v.text.strip():
                    page = v.prov.page + 1 if v.prov else (
                        calc.pages[0] + 1 if calc.pages else None)
                    said.append(f"'{v.text.strip()}'"
                                + (f" (PDF page {page})" if page else ""))
                    if page:
                        pages.append(page)
        if not said:
            return _not_given(f"no calculation on the record prints a "
                              f"{words}")
        return _given("the calculations print: "
                      + "; ".join(list(dict.fromkeys(said))[:MAX_LISTED]),
                      pages, "code's reading off the calculation printouts")
    return read


#: The chemical results a sample's line lists, in order.
_CHEMICAL = ("pH", "resistivity", "resistivity_minimum", "sulfate",
             "chloride", "sulfides", "redox")


#: The checkbox marks Document Intelligence leaves in a cell's text.
_RE_CHECKBOX = re.compile(r":(?:un)?selected:")


def _reported(value: Any) -> str:
    if hasattr(value, "unit"):
        said = _depth(value)
    else:
        said = f"{value:g}" if isinstance(value, float) else str(value)
    return " ".join(_RE_CHECKBOX.sub(" ", said).split())


def _chemical(src: _Source, q: Question) -> Answer:
    rows: List[str] = []
    pages: List[int] = []

    def add(hole: str, sample: str, depth: Any, values: Any,
            where: Sequence[int]) -> None:
        read = {name: _reported(getattr(values, name)) for name in _CHEMICAL
                if getattr(values, name, None) not in (None, "")}
        got = [f"{name.replace('_', ' ')} {said}"
               for name, said in read.items() if said]
        if not got:
            return
        label = " ".join(s for s in (hole, sample) if s) or "a sample"
        at = f" at {_depth(depth)}" if depth is not None else ""
        rows.append(f"{label}{at}: " + ", ".join(got))
        pages.extend(p + 1 for p in where[:1])

    for test in src.record.lab_tests:
        if test.result is None:
            continue
        if test.kind == "chemical":
            add(test.investigation_id, test.sample_id, test.depth_top,
                test.result, test.pages)
        elif test.kind == "summary_table":
            for row in test.result.rows:
                add(row.investigation_id, row.sample_id, row.depth_top, row,
                    test.pages)
    if not rows:
        return _not_given("no chemical test result on the record")
    more = len(rows) - MAX_LISTED
    return _given("the laboratory's chemical results: "
                  + "; ".join(rows[:MAX_LISTED])
                  + (f"; and {more} more sample(s)" if more > 0 else ""),
                  pages, "code's reading off the laboratory sheets, to check "
                         "the summary against")


def _has_volume(src: _Source, q: Question) -> Answer:
    vols = [d for d in src.record.bound_documents if d.kind == "volume"]
    if not vols:
        return _not_given("no bound volume found; code cannot show a no")
    return _given(True, [d.first_page + 1 for d in vols],
                  "a bound volume on PDF pages " + ", ".join(
                      f"{d.first_page + 1}-{d.last_page + 1}" for d in vols))


def _no_argument(read: Callable[[_Source, Question], Answer]
                 ) -> Callable[[str], Callable[[_Source, Question], Answer]]:
    """A reading that takes no argument."""
    def make(arg: str) -> Callable[[_Source, Question], Answer]:
        if arg:
            raise ValueError(f"this reading takes no argument (not {arg!r})")
        return read
    return make


#: CODE'S READINGS of a record, by the name a code map gives them ("name"
#: or "name:argument"). Each makes the reading from its argument, and
#: raises ValueError on one it cannot use. What each reads:
#:
#: * ``bound_documents``: the bound reports (title, firm, date, kind, PDF
#:   pages), one record each, in the row's parts.
#: * ``has_bound_volume``: yes when a bound volume is found.
#: * ``pages_labelled:<label>``: the PDF pages with that page label.
#: * ``page_labels``: every page's label (its class) and printed number.
#: * ``hole_count:<log kinds>``: holes logged, of those kinds, in this file
#:   and its bound volumes; ``prior_hole_count:<log kinds>``: in its bound
#:   prior reports.
#: * ``bound_prior_count``: how many prior reports are bound in.
#: * ``caption_count:<tables or figures>``: captions in the main body.
#: * ``soft_layers``: yes when a log layer is described as soft.
#: * ``water_readings``: the logs' water readings, as text to check against.
#: * ``calculation_of:<calculation kinds>``: yes when a printout of those
#:   kinds is on the record; ``seismic_case:<kinds>``: when one names a
#:   seismic case.
#: * ``calculation_list``: the printouts by kind, as text to check against.
#: * ``printed_value:<name|name>``: what the printouts print under a name.
#: * ``chemical_results``: the laboratory's chemical results, as text.
#:
#: A yes-or-nothing reading never answers no: no log or no calculation on
#: the record is not a no, because a report can describe what it does not
#: bind in.
READINGS: Dict[str, Callable[[str], Callable[[_Source, Question], Answer]]] = {
    "bound_documents": _no_argument(_bound_documents),
    "has_bound_volume": _no_argument(_has_volume),
    "pages_labelled": _pages_labelled,
    "page_labels": _no_argument(_page_labels),
    "hole_count": _count(prior=False),
    "prior_hole_count": _count(prior=True),
    "bound_prior_count": _no_argument(_bound_priors),
    "caption_count": _captions,
    "soft_layers": _no_argument(_soft),
    "water_readings": _no_argument(_water),
    "calculation_of": _has_calc,
    "seismic_case": _seismic,
    "calculation_list": _no_argument(_calc_list),
    "printed_value": _printed_value,
    "chemical_results": _no_argument(_chemical),
}


def code_answers(record: Any, questions: Sequence[Question], *,
                 bound: Optional[Mapping[str, Any]] = None,
                 printed_pages: Optional[Mapping[int, Any]] = None
                 ) -> Dict[str, Answer]:
    """Code's answer to each row answered by "Code" or "Both, compared",
    from a finished report record (a ``ReportRecord`` or its JSON), by the
    list's code map (the ``readings`` of a :class:`QuestionList`).

    ``bound`` holds the records of the reports bound in it, by bound_id: a
    bound volume's logs count towards this investigation, a prior report's
    towards the prior counts. ``printed_pages`` maps a 0-based page to its
    printed page number, which the record does not keep.

    Never a made-up value: a row code cannot answer is ``not said`` with
    the reason in its note ("not on the record" when the code map gives it
    no reading)."""
    src = _Source(_as_record(record), dict(bound or {}),
                  {int(k): str(v) for k, v in (printed_pages or {}).items()
                   if v not in (None, "")})
    readings = getattr(questions, "readings", None) or {}
    mapped = bool(getattr(questions, "map_name", "") or readings)
    out: Dict[str, Answer] = {}
    for q in questions:
        if q.answered_by == "model":
            continue
        if q.key not in readings:
            out[q.key] = _not_given("not on the record" if mapped else
                                    "not read: the list has no code map")
            continue
        name, arg = readings[q.key]
        out[q.key] = _fit(q, READINGS[name](arg)(src, q))
    return out


# ---------------------------------------------------------------------------
# a Both row: code's reading beside the model's
# ---------------------------------------------------------------------------

def _norm(value: Any) -> str:
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, Mapping):
        return " | ".join(_norm(v) for v in value.values())
    text = " ".join(str(value).lower().split())
    return text.strip(" .,;:")


def _same(question: Question, a: Any, b: Any) -> Optional[bool]:
    """Whether two answers to ``question`` agree; None when no rule compares
    them (a written description beside code's list)."""
    kind = question.data_type
    try:
        if kind == "integer":
            return int(a) == int(b)
        if kind == "boolean":
            return bool(a) == bool(b)
        if kind == "integer or list of integers":
            # Page numbers, mostly: the same pages in any order agree.
            return {int(x) for x in a} == {int(x) for x in b}
    except (TypeError, ValueError):
        return False
    if kind in ("list of records", "dictionary") and question.parts:
        # Records agree on their first part (a Dictionary's key): the same
        # bound reports by title, the same pages by page. The other parts
        # are read differently by the two (a page span, a firm's spelling),
        # and the owner compares them on the sheet.
        first = question.parts[0].name

        def firsts(records: Any) -> set:
            return {_norm(r.get(first) if isinstance(r, Mapping) else r)
                    for r in records or ()}
        return firsts(a) == firsts(b)
    if kind in ("list of strings", "list of records", "dictionary"):
        return {_norm(x) for x in a or ()} == {_norm(x) for x in b or ()}
    if kind == "date" or question.options:
        return _norm(a) == _norm(b)
    return None


def compare(question: Question, model: Optional[Answer],
            code: Optional[Answer]) -> Answer:
    """A "Both" row: the model's answer, with code's reading kept beside it
    (``code_value`` and the note) and the From column saying whether the
    two agree. Counts compare as whole numbers, lists as sets of normalised
    text, a closed-list or date answer as normalised text; a written
    description is compared by eye. Code with no reading leaves the model's
    answer as it is, with code's reason in the note."""
    model = model or Answer(status="", note="no answer came back")
    out = Answer(value=model.value, status=model.status,
                 pdf_pages=list(model.pdf_pages), note=model.note,
                 source=FROM_TEXT)
    if code is None or code.status != "answered":
        return out.noted(f"code: {code.note}" if code else "")
    out.code_value = code.value
    shown = _shown(question, code.value)
    if not model.status:                          # no answer came back
        out.source = ""
        return out.noted(f"code: {shown}")
    if model.status != "answered":
        out.source = FROM_DISAGREE
        return out.noted(f"code: {shown}")
    same = _same(question, model.value, code.value)
    out.source = (FROM_BY_EYE if same is None
                  else FROM_AGREE if same else FROM_DISAGREE)
    return out.noted(f"code: {shown}")


# ---------------------------------------------------------------------------
# the calls
# ---------------------------------------------------------------------------

def pages_to_read(record: Any) -> List[int]:
    """The 0-based pages the calls read, by today's narrative reader's rule:
    every page labelled cover, letter, contents or narrative, and the front
    of the report whatever it was labelled."""
    from report_ingest.narrative_reader import FRONT_LABELS, FRONT_PAGES

    record = _as_record(record)
    n = int(record.document.n_pages or len(record.page_labels))
    wanted = {lab.page for lab in record.page_labels
              if lab.label in FRONT_LABELS}
    wanted.update(range(min(FRONT_PAGES, n) if n else FRONT_PAGES))
    return sorted(p for p in wanted if not n or p < n)


def page_texts(doc: Any, pages: Sequence[int]) -> Dict[int, str]:
    """Each page's text with its tables, as the narrative reader serialises
    it (:func:`report_ingest.narrative_reader.serialise_page`), without its
    0-based heading: the calls head each page with its PDF page."""
    from report_ingest.narrative_reader import serialise_page

    out: Dict[int, str] = {}
    for page in pages:
        text = serialise_page(doc, int(page))
        out[int(page)] = text.split("\n", 1)[1] if "\n" in text else ""
    return out


def page_list_text(record: Any,
                   printed_pages: Optional[Mapping[int, Any]] = None) -> str:
    """What each page of the file is, one line per run of pages with one
    class: "PDF pages 4-9 (printed 1-6): narrative"."""
    record = _as_record(record)
    printed = {int(k): str(v) for k, v in (printed_pages or {}).items()
               if v not in (None, "")}
    runs: List[List[Any]] = []
    for lab in sorted(record.page_labels, key=lambda lab: lab.page):
        if runs and runs[-1][1] == lab.page - 1 and runs[-1][2] == lab.label:
            runs[-1][1] = lab.page
        else:
            runs.append([lab.page, lab.page, lab.label])
    lines = []
    for first, last, label in runs:
        where = (f"PDF page {first + 1}" if first == last
                 else f"PDF pages {first + 1}-{last + 1}")
        shown = [printed[p] for p in (first, last) if p in printed]
        if shown:
            where += (f" (printed {shown[0]})" if len(set(shown)) == 1
                      else f" (printed {shown[0]}-{shown[-1]})")
        lines.append(f"{where}: {label.replace('_', ' ')}")
    return "\n".join(lines)


def _shown(question: Optional[Question], value: Any) -> str:
    """A value as one readable cell."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, list):
        if all(isinstance(v, int) and not isinstance(v, bool) for v in value):
            return _span(value)
        if all(isinstance(v, Mapping) for v in value):
            return "\n".join(_record_text(question, v) for v in value)
        return "; ".join(_shown(question, v) for v in value)
    return str(value)


def _record_text(question: Optional[Question], record: Mapping[str, Any]
                 ) -> str:
    roles = {p.name: p.role for p in (question.parts if question else ())}
    key = next((n for n, r in roles.items() if r == "key"), None)
    value = next((n for n, r in roles.items() if r == "value"), None)
    if key and value and set(record) <= {key, value}:
        return f"{_shown(None, record.get(key))}: " \
               f"{_shown(None, record.get(value))}"
    return "; ".join(f"{name.replace('_', ' ')}: "
                     f"{' '.join(_shown(None, v).split())}"
                     for name, v in record.items()
                     if v not in (None, "", []))


_HOW_TO_ANSWER = """\
HOW TO ANSWER. Answer every question in the answer schema. Each answer has:
- status: "answered" when the pages give the answer; "not said" when the \
report is silent on it; "does not apply" when the question has no bearing on \
a report like this one. Keep the three apart. The value stays null unless the \
status is "answered".
- pdf_pages: the pages the answer comes from, numbered as the headings above \
number them (1 for the file's first page).
- note: anything a reader of the answer needs to know; empty if nothing."""


def _instruction(record: Any, report_id: str, rules: Sequence[str],
                 pages_list: str, screening: Mapping[str, Answer],
                 by_key: Mapping[str, Question], first: bool) -> str:
    # Code's readings are never sent: on a Both row the model's answer and
    # code's must be two readings made apart for their agreeing to mean
    # anything. Code's stay on the answer sheet for the owner.
    parts = [f"The pages above are from report {report_id}, a file of "
             f"{record.document.n_pages} PDF pages. "
             + ("Answer the screening questions in the answer schema."
                if first else
                "Answer the questions in the answer schema.")]
    if rules:
        parts.append("THE RULES FOR EVERY ANSWER\n"
                     + "\n".join(f"- {r}" for r in rules))
    parts.append(_HOW_TO_ANSWER)
    parts.append("WHAT EACH PAGE OF THE FILE IS (the ingest's page classes, "
                 "for questions about pages the text above does not show)\n"
                 + (pages_list or "(no page classes on the record)"))
    if screening:
        said = [f"- {key}: " + (_shown(by_key.get(key), a.value)
                                if a.status == "answered"
                                else a.status or "no answer")
                for key, a in screening.items()]
        parts.append("THE SCREENING ANSWERS ALREADY GIVEN\n" + "\n".join(said))
    return "\n\n".join(parts)


def _pages_block(texts: Mapping[int, str], printed: Mapping[int, str]
                 ) -> str:
    out = []
    for page in sorted(texts):
        head = f"=== PDF page {page + 1}"
        if printed.get(page):
            head += f" (printed {printed[page]})"
        out.append(f"{head} ===\n{texts[page].strip()}")
    return "\n\n".join(out) if out else "(no page text was given)"


@dataclass
class QuestionRun:
    """What :func:`run_questions` returns: every row's answer, the calls,
    the rows screening closed (key -> why) and any warnings."""

    report_id: str
    answers: Dict[str, Answer]
    calls: List[Dict[str, Any]] = field(default_factory=list)
    closed: Dict[str, str] = field(default_factory=dict)
    warnings: List[str] = field(default_factory=list)
    pages: List[int] = field(default_factory=list)
    prompt_version: str = PROMPT_VERSION


def _from_reply(reply: Any, schema: type) -> Optional[Any]:
    parsed = getattr(reply, "parsed", None)
    if parsed is None or getattr(reply, "cut_off", False):
        return None
    if isinstance(parsed, schema):
        return parsed
    data = parsed.model_dump() if isinstance(parsed, BaseModel) else parsed
    return schema.model_validate(data)


def _checked(question: Question, answer: Answer, n_pages: int) -> Answer:
    """A model's answer with what does not hold up noted, never dropped
    without a word: a page outside the file, a value given with another
    status, a blank on a row the list does not expect blank."""
    if n_pages:
        outside = [p for p in answer.pdf_pages if not 1 <= p <= n_pages]
        if outside:
            answer.pdf_pages = [p for p in answer.pdf_pages
                                if 1 <= p <= n_pages]
            answer.noted(f"page(s) {', '.join(map(str, outside))} are not "
                         f"in the {n_pages}-page file and were left out")
    if answer.status != "answered" and answer.value not in (None, [], ""):
        answer.noted(f"a value came with status '{answer.status}': "
                     f"{_shown(question, answer.value)}")
        answer.value = None
    if answer.status == "answered" and answer.value is None:
        answer.noted("answered with no value")
    if answer.status == "not said" and not question.allow_null:
        answer.noted("the list does not expect this row blank")
    return answer


def _ask(engine: Any, call: str, asked: Sequence[Question], text: str,
         record: Any, run: QuestionRun) -> Dict[str, Answer]:
    """One call: the text, the answer schema of ``asked``. Its answers by
    key; none when no structured answer came back (said in a warning)."""
    from report_ingest.engine import text_block, user

    if not asked:
        return {}
    schema = answer_schema(asked, "ScreeningAnswers" if call == "screening"
                           else "QuestionAnswers")
    reply = engine.complete([user(text_block(text))], output_format=schema)
    usage = getattr(reply, "usage", None)
    run.calls.append({
        "call": call, "model": getattr(reply, "model", "")
        or str(getattr(engine, "name", "")),
        "questions": len(asked), "stop_reason": getattr(reply, "stop_reason",
                                                        ""),
        "input_tokens": getattr(usage, "input_tokens", 0),
        "output_tokens": getattr(usage, "output_tokens", 0),
        "chars_sent": len(text)})
    try:
        parsed = _from_reply(reply, schema)
    except ValidationError as exc:
        parsed = None
        run.warnings.append(f"the {call} call's answer did not fit the "
                            f"schema: {str(exc)[:300]}")
    run.calls[-1]["answered"] = parsed is not None
    if parsed is None:
        run.warnings.append(
            f"the {call} call gave no structured answer (stop_reason "
            f"{getattr(reply, 'stop_reason', '')!r})")
        return {}
    n_pages = int(record.document.n_pages or 0)
    out = {}
    for q in asked:
        got = getattr(parsed, q.key)
        out[q.key] = _checked(q, Answer(
            value=_plain(got.value), status=got.status,
            pdf_pages=list(got.pdf_pages), note=got.note.strip(),
            source=FROM_TEXT), n_pages)
    return out


def closed_rows(questions: Sequence[Question],
                screening: Mapping[str, Answer]) -> Dict[str, str]:
    """The rows the screening answers close by the list's gates (its code
    map's :class:`Gate` entries), each with why: key -> "screening answered
    ..."."""
    by_key = {q.key: q for q in questions}
    out: Dict[str, str] = {}
    for gate in getattr(questions, "gates", None) or ():
        answer = screening.get(gate.key)
        if gate.key not in by_key or answer is None \
                or answer.status != "answered" \
                or answer.value is not gate.closes_when:
            continue
        why = (f"screening answered {by_key[gate.key].output or gate.key} "
               f"{'yes' if gate.closes_when else 'no'}")
        for q in questions:
            if q.screening or q.key in out:
                continue
            if gate.category == "*" \
                    or q.category.lower() == gate.category.lower():
                out[q.key] = why
    return out


def _flag_both_yes(answers: Mapping[str, Answer],
                   pairs: Sequence[Tuple[str, str]]) -> List[str]:
    """Flag, on both answers, each pair that cannot both be yes and was;
    the pairs flagged, one line each."""
    flagged = []
    for a, b in pairs:
        if all(answers.get(k) is not None and answers[k].status == "answered"
               and answers[k].value is True for k in (a, b)):
            for k in (a, b):
                answers[k].noted(f"code flags {a} and {b} both answered yes")
            flagged.append(f"{a} and {b}")
    return flagged


def run_questions(record: Any, pages: Mapping[int, str],
                  questions: Sequence[Question], engine: Any, *,
                  bound: Optional[Mapping[str, Any]] = None,
                  printed_pages: Optional[Mapping[int, Any]] = None,
                  report_id: str = "") -> QuestionRun:
    """The list asked of one report: a screening call, then one call for
    the rest.

    ``pages`` is the text of the pages to read (0-based page -> text, as
    :func:`page_texts` gives it over :func:`pages_to_read`), kept to today's
    narrative input budget with the front of the report first. The
    screening call reads the first :data:`SCREENING_PAGES` of them; the main
    call reads them all. Both get the Rules sheet and the page list
    (:func:`page_list_text`), and the main call the screening answers.
    Neither gets code's readings: a Both row's model answer is read apart
    from code's, which :func:`compare` sets beside it afterwards.
    ``engine`` is any :class:`~report_ingest.engine.Engine`; each call
    offers its answer schema as the structured output, as the narrative
    stage does.

    A row answered by Code is not asked. A row screening closes
    (:func:`closed_rows`) is not asked either and is ``does not apply``,
    with code's reading kept in its note where it has one. Screening that
    contradicts itself (a pair that cannot both be yes, both yes) closes
    nothing, with a warning: a wrong N/A would lose answers. A call that
    gives no structured answer leaves its rows with status "" and a
    warning; when screening gives none, nothing is closed."""
    from report_ingest.narrative_reader import (CHARS_PER_TOKEN, FRONT_PAGES,
                                                INPUT_TOKEN_BUDGET,
                                                _within_budget)

    record = _as_record(record)
    rid = report_id or record.document.report_id
    rules = list(getattr(questions, "rules", None) or [])
    printed = {int(k): str(v) for k, v in (printed_pages or {}).items()
               if v not in (None, "")}
    code = code_answers(record, questions, bound=bound, printed_pages=printed)
    texts = {int(p): str(t) for p, t in pages.items()}
    kept, dropped = _within_budget(sorted(texts), texts, FRONT_PAGES,
                                   INPUT_TOKEN_BUDGET * CHARS_PER_TOKEN)
    run = QuestionRun(report_id=rid, answers={}, pages=kept)
    if dropped:
        run.warnings.append(
            f"{len(dropped)} page(s) were over the "
            f"{INPUT_TOKEN_BUDGET:,}-token input budget and were not sent: "
            f"PDF pages {_span([p + 1 for p in dropped])}")
    pages_list = page_list_text(record, printed)
    by_key = {q.key: q for q in questions}

    asked = [q for q in screening_questions(questions)
             if q.answered_by != "code"]
    front = {p: texts[p] for p in kept[:SCREENING_PAGES]}
    got = _ask(engine, "screening", asked, _pages_block(front, printed)
               + "\n\n" + _instruction(record, rid, rules, pages_list, {},
                                       by_key, first=True),
               record, run)
    contradicted = _flag_both_yes(
        got, getattr(questions, "not_both_yes", None) or ())
    if contradicted:
        run.warnings.append(
            f"screening answered {'; '.join(contradicted)} both yes, which "
            f"cannot both be; nothing is closed and every row is asked")
    run.closed = closed_rows(questions, got) if got and not contradicted \
        else {}
    asked = [q for q in main_questions(questions)
             if q.answered_by != "code" and q.key not in run.closed]
    got.update(_ask(engine, "main", asked, _pages_block(
        {p: texts[p] for p in kept}, printed) + "\n\n" + _instruction(
            record, rid, rules, pages_list,
            {k: v for k, v in got.items() if by_key[k].screening}, by_key,
            first=False), record, run))

    for q in questions:
        if q.key in run.closed:
            answer = Answer(status="does not apply", source=FROM_SCREENING,
                            note=run.closed[q.key])
            if q.key in code and code[q.key].status == "answered":
                answer.noted(f"code's reading: "
                             f"{_shown(q, code[q.key].value)}")
            run.answers[q.key] = answer
        elif q.answered_by == "code":
            run.answers[q.key] = code[q.key]
        elif q.answered_by == "both":
            run.answers[q.key] = compare(q, got.get(q.key), code.get(q.key))
        else:
            run.answers[q.key] = got.get(q.key) or Answer(
                status="", note="no answer came back for this row")
    return run


# ---------------------------------------------------------------------------
# the answer sheet
# ---------------------------------------------------------------------------

#: The answer sheet's columns, in order.
SHEET_COLUMNS: Tuple[str, ...] = (
    "Output", "Category", "Prompt", "Key", "Answer", "PDF pages", "From",
    "Note", "Right?", "Correct answer")


def _answer_cell(question: Question, answer: Answer) -> str:
    """The Answer cell: empty where the report is silent, N/A where the
    question has no bearing."""
    if answer.status == "does not apply":
        return "N/A"
    if answer.status != "answered":
        return ""
    if answer.value == []:
        return "none"
    return _shown(question, answer.value)


def _tab_name(text: str, taken: Sequence[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", text).strip()[:31] or "Answer"
    base, n = name, 2
    while name in taken:
        name = f"{base[:28]} {n}"
        n += 1
    return name


def answer_sheet(questions: Sequence[Question], answers: Mapping[str, Answer],
                 path: Union[str, Path], *, report_id: str = "",
                 calls: Sequence[Mapping[str, Any]] = (),
                 prompt_version: str = PROMPT_VERSION) -> Path:
    """The answers written to ``path`` as an xlsx in the owner's layout
    (:data:`SHEET_COLUMNS`), one row per question in the list's order, with
    empty "Right?" and "Correct answer" columns for the owner. Lists and
    records are written readably in one cell; a list of more than
    :data:`MAX_CELL_RECORDS` records goes on its own tab, named by its
    Output. A "Run" tab gives the report, the date, the list's file name and
    modified time, its code map's file name, and the model of each call."""
    import openpyxl
    from openpyxl.styles import Alignment, Font

    path = Path(path)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Answers"
    sheet.append(list(SHEET_COLUMNS))
    tabs: List[Tuple[str, Question, List[Mapping[str, Any]]]] = []
    for q in questions:
        answer = answers.get(q.key) or Answer(status="", note="not run")
        cell = _answer_cell(q, answer)
        value = answer.value
        if (answer.status == "answered" and isinstance(value, list)
                and len(value) > MAX_CELL_RECORDS
                and all(isinstance(v, Mapping) for v in value)):
            name = _tab_name(q.output or q.key,
                             [t[0] for t in tabs] + ["Answers", "Run"])
            tabs.append((name, q, value))
            cell = f"{len(value)} rows, on the '{name}' tab"
        sheet.append([q.output, q.category, q.prompt, q.key, cell,
                      _span(answer.pdf_pages), answer.source, answer.note,
                      None, None])
    for column, width in zip("ABCDEFGHIJ",
                             (24, 16, 60, 26, 50, 14, 20, 50, 10, 30)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    for name, q, records in tabs:
        tab = book.create_sheet(name)
        tab.append([p.name for p in q.parts])
        for record in records:
            tab.append([_shown(q, record.get(p.name)) if isinstance(
                record.get(p.name), list) else record.get(p.name)
                for p in q.parts])
        for cell in tab[1]:
            cell.font = Font(bold=True)
        tab.freeze_panes = "A2"

    info = book.create_sheet("Run")
    modified = getattr(questions, "modified", None)
    info.append(["Report", report_id])
    info.append(["Run date", datetime.now().strftime("%Y-%m-%d %H:%M")])
    info.append(["Question list", getattr(questions, "name", "")])
    info.append(["List modified", modified.strftime("%Y-%m-%d %H:%M")
                 if modified else ""])
    info.append(["Code map", getattr(questions, "map_name", "")
                 or "(none: code read no row)"])
    info.append(["Prompt version", prompt_version])
    info.append([])
    info.append(["Call", "Model", "Questions", "Answered", "Stop reason",
                 "Input tokens", "Output tokens"])
    for call in calls:
        info.append([call.get("call"), call.get("model"),
                     call.get("questions"), call.get("answered"),
                     call.get("stop_reason"), call.get("input_tokens"),
                     call.get("output_tokens")])
    if not calls:
        info.append(["(no model call: code's rows only)"])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 40
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path
